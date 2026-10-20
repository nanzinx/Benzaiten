"""
Testes do relatório da URA (relatorio_ura.py) e das conversões do leitura.py.

Todos os nomes, valores e CPFs são INVENTADOS e criados dentro do teste.
"""

from datetime import date

import openpyxl
import pytest

import relatorio_ura as ru
from leitura import converter_data, converter_valor_br, tempo_para_segundos

DIA = date(2026, 10, 19)


def escrever_csv(caminho, linhas, codificacao="cp1252"):
    texto = "\r\n".join(";".join(f'"{c}"' for c in linha) for linha in linhas) + "\r\n"
    caminho.write_text(texto, encoding=codificacao)


# --- Conversões --------------------------------------------------------------

@pytest.mark.parametrize("texto, esperado", [
    ("R$ 1.234,56", 1234.56), ("10,50", 10.5), ("1,3", 1.3), ("633,0333", 633.0333),
    ("1.500", 1500.0), ("2.345.678,90", 2345678.9), ("1500.00", 1500.0), ("", 0.0),
    ("R$ -", 0.0), ("-2,5", -2.5), (7, 7.0),
])
def test_converter_valor_br(texto, esperado):
    assert converter_valor_br(texto) == pytest.approx(esperado)


def test_converter_tempo_e_data():
    assert tempo_para_segundos("00:01:04") == 64
    assert tempo_para_segundos("01:02:05") == 3725
    assert tempo_para_segundos("") == 0
    assert converter_data("19/10/2026 14:35:00") == DIA
    assert converter_data("2026-10-19") == DIA
    assert converter_data("- TOTAL GERAL -") is None


# --- Cadastro ----------------------------------------------------------------

def test_equipe_e_o_primeiro_nome_do_supervisor():
    assert ru.limpar_supervisor("(1) MARIA DE FÁTIMA SOUZA, (2) OUTRO") == "MARIA DE FÁTIMA SOUZA"
    assert ru.nome_da_equipe("(1) MARIA DE FÁTIMA SOUZA") == "Maria"


def test_atualizar_cadastro_mantem_os_apelidos(tmp_path):
    arquivo = tmp_path / "operadores.xlsx"
    ru.atualizar_cadastro([{"NOME": "ANA", "SUPERVISOR(ES)": "(1) SUP UM", "STATUS": "ATIVO"}], arquivo)
    wb = openpyxl.load_workbook(arquivo)
    wb["Apelidos"].append(["ANNA", "ANA"])
    wb.save(arquivo)

    ru.atualizar_cadastro([{"NOME": "ANA", "SUPERVISOR(ES)": "(1) SUP UM", "STATUS": "ATIVO"},
                           {"NOME": "BETO", "SUPERVISOR(ES)": "(1) SUP DOIS", "STATUS": "INATIVO"}], arquivo)
    ativos, inativos, apelidos = ru.ler_cadastro(arquivo)

    assert ativos == {"ANA": ("ANA", "Sup")}
    assert inativos == {"BETO"}
    assert apelidos == {"ANNA": "ANA"}  # a linha "(exemplo)" é ignorada


# --- Cenário completo com dados fictícios -------------------------------------

@pytest.fixture
def entrada(tmp_path):
    """
    Pasta entrada/ fictícia:
      ANA e BRUNO (equipe Sup), CARLA (equipe Chefe), DIEGO inativo (ex-colaborador).
      Na tarifação, CARLA aparece como "KARLA" (ligada pela aba Apelidos).
    """
    pasta = tmp_path / "entrada"
    pasta.mkdir()
    escrever_csv(pasta / "OPERADORES.csv", [
        ["", "Cód. Usuário", "? Nome", "CPF", "Supervisor(es)", "Status"],
        ["", "1", "ANA ALVES", "", "(1) SUP UM", "ATIVO"],
        ["", "2", "BRUNO BRITO", "", "(1) SUP UM", "ATIVO"],
        ["", "3", "CARLA CRUZ", "", "(1) CHEFE DOIS", "ATIVO"],
        ["", "4", "DIEGO DIAS", "", "(1) CHEFE DOIS", "INATIVO"],
    ])
    escrever_csv(pasta / "tme_tma.csv", [
        ["Data", "Operador", "Supervisor", "TME", "TMA"],
        ["19/10/2026", "ANA ALVES", "SUP UM", "00:01:00", "00:02:00"],
        ["19/10/2026", "BRUNO BRITO", "SUP UM", "00:00:30", "00:00:40"],
        ["19/10/2026", "CARLA CRUZ", "CHEFE DOIS", "00:00:20", "00:00:50"],
        ["", "- TOTAL GERAL -", "SUP UM", "00:00:37", "00:01:10"],
    ])
    escrever_csv(pasta / "tme_tma_antigo.csv", [  # de outro dia: deve ser ignorado
        ["Data", "Operador", "Supervisor", "TME", "TMA"],
        ["16/10/2026", "ANA ALVES", "SUP UM", "09:00:00", "09:00:00"],
    ])
    escrever_csv(pasta / "leads.csv", [
        ["Data", "Operador", "Supervisor", "Leads. Atend."],
        ["19/10/2026", "ANA ALVES", "SUP UM", "30"],
        ["19/10/2026", "BRUNO BRITO", "SUP UM", "10"],
        ["19/10/2026", "CARLA CRUZ", "CHEFE DOIS", ""],
        ["", "- TOTAL GERAL -", "SUP UM", "40"],
    ])
    escrever_csv(pasta / "VENDITORE.csv", [
        ["Descrição", "Valor R$", "Qtde. Lig", "Minutos"],
        ["<LIGAÇÕES SEM DERIVAÇÃO P/ OPERADOR>", "100,00", "500", "10"],
        ["ANA ALVES", "1,50", "10", "5"],
        ["OPER VIRTUAL 1", "30,00", "100", "50"],
        ["OPERADOR VIRTUAL - DTMF", "10,00", "50", "20"],
        ["DIEGO DIAS", "2,00", "5", "1"],          # ex-colaborador
        ["ZECA FANTASMA", "0,70", "3", "1"],       # não existe no cadastro
    ])
    escrever_csv(pasta / "4NET.csv", [
        ["Descrição", "Valor R$", "Qtde. Lig", "Minutos"],
        ["<LIGAÇÕES SEM DERIVAÇÃO P/ OPERADOR>", "20,50", "100", "5"],
        ["BRUNO BRITO", "1.000,25", "10", "5"],
        ["KARLA CRUZ", "3,00", "10", "5"],         # grafia diferente: ligada por Apelidos
    ])
    escrever_csv(pasta / "esteira.csv", [
        ["Proposta", "Data/Hora", "Etapa", "CPF", "Agente", "Valor Contrato"],
        ["", "19/10/2026 09:00:00", "Pago", "CLIENTE-000001", "CARLA CRUZ", "5000,00"],
        ["", "19/10/2026 10:00:00", "Andamento", "CLIENTE-000001", "CARLA CRUZ", "1000,00"],  # mesmo CPF
        ["", "19/10/2026 11:00:00", "Pendente", "CLIENTE-000002", "CARLA CRUZ", "500,00"],   # não é cliente
        ["", "19/10/2026 12:00:00", "Cancelado", "CLIENTE-000003", "CARLA CRUZ", "9999,00"],  # não conta
        ["", "19/10/2026 13:00:00", "Pago", "CLIENTE-000004", "BRUNO BRITO", "6500,00"],
        ["", "18/10/2026 13:00:00", "Pago", "CLIENTE-000005", "ANA ALVES", "7777,00"],     # outro dia
        ["", "19/10/2026 14:00:00", "Pago", "CLIENTE-000006", "DIEGO DIAS", "300,00"],     # ex-colaborador
    ], "utf-8-sig")
    (pasta / "anotacoes.txt").write_text("não é exportação", encoding="utf-8")
    return pasta


@pytest.fixture
def resultado(entrada, tmp_path):
    cadastro = tmp_path / "config" / "operadores.xlsx"
    entradas, _ = ru.ler_entradas(entrada)
    ru.atualizar_cadastro(entradas["operadores"][0][1], cadastro)
    wb = openpyxl.load_workbook(cadastro)
    wb["Apelidos"].append(["KARLA CRUZ", "CARLA CRUZ"])
    wb.save(cadastro)
    ativos, inativos, apelidos = ru.ler_cadastro(cadastro)
    return ru.calcular(DIA, entradas, ativos, inativos, apelidos)


def test_arquivos_sao_reconhecidos_pelas_colunas(entrada):
    entradas, desconhecidos = ru.ler_entradas(entrada)
    assert {t: sorted(a.name for a, _ in v) for t, v in entradas.items()} == {
        "esteira": ["esteira.csv"], "operadores": ["OPERADORES.csv"],
        "tarifacao": ["4NET.csv", "VENDITORE.csv"],
        "desempenho": ["tme_tma.csv", "tme_tma_antigo.csv"], "leads": ["leads.csv"],
    }
    assert desconhecidos == []  # .txt nem é considerado


def test_tma_tmo_e_leads_so_do_dia(resultado):
    ana = resultado.operadores["ANA ALVES"]
    assert (ana.tma, ana.tmo, ana.leads) == (120, 60, 30)  # o arquivo de 16/10 foi ignorado
    assert any("IGNORADO: é de 16/10/2026" in obs for _, _, obs in resultado.arquivos)


def test_custos_das_duas_tarifacoes_e_rateio_dos_robos(resultado):
    ops = resultado.operadores
    assert resultado.custo_robos == pytest.approx(40.0)          # OPER VIRT + OPERADOR VIRT
    assert resultado.sem_derivacao == pytest.approx(120.5)       # 100,00 + 20,50
    # Robôs (R$ 40) divididos pelos leads: ANA 30/40, BRUNO 10/40, CARLA 0
    assert ops["ANA ALVES"].custo == pytest.approx(1.5 + 30.0)
    assert ops["BRUNO BRITO"].custo == pytest.approx(1000.25 + 10.0)
    assert ops["CARLA CRUZ"].custo == pytest.approx(3.0)         # veio como KARLA (Apelidos)


def test_ex_colaborador_e_nome_desconhecido_sao_ignorados_e_avisados(resultado):
    assert "DIEGO DIAS" not in resultado.operadores
    assert resultado.ex_colaboradores == {"DIEGO DIAS": {"Tarifação", "Vanguard"}}
    assert resultado.sem_vinculo == {"ZECA FANTASMA": {"Tarifação"}}
    assert resultado.custo_ignorado == pytest.approx(2.70)
    assert resultado.vendas_ignoradas == pytest.approx(300.0)


def test_vendas_e_clientes_efetivados(resultado):
    carla = resultado.operadores["CARLA CRUZ"]
    assert carla.vendas == pytest.approx(6500.0)  # Pago + Andamento + Pendente; Cancelado fora
    assert carla.clientes == 1                    # CPF repetido conta uma vez; Pendente não conta
    assert resultado.operadores["ANA ALVES"].vendas == 0  # venda de 18/10 não entra


def test_ordem_por_vendas_desempate_por_leads_e_totais(resultado):
    linhas = ru.linhas_do_relatorio(resultado)
    nomes = [l["nome"] for l in linhas]
    # CARLA e BRUNO venderam R$ 6.500; BRUNO tem mais leads (10 x 0). ANA não vendeu.
    assert nomes[:3] == ["BRUNO BRITO", "CARLA CRUZ", "ANA ALVES"]
    assert nomes[3] == ""  # separador amarelo
    # Equipes por vendas: Sup (6.500) e Chefe (6.500) empatam -> ordem alfabética
    assert nomes[4:6] == ["TOTAL EQUIPE - Chefe", "TOTAL EQUIPE - Sup"]
    assert nomes[6:] == ["TOTAL OPERADORES", "LIGAÇÕES SEM DERIVAÇÃO P/ OPERADOR", "TOTAL GERAL (c/ sem derivação)"]

    sup = linhas[5]
    assert (sup["tma"], sup["tmo"], sup["leads"]) == (160, 90, 40)  # TMA/TMO da equipe = soma

    total_op, sem_der, total = linhas[6], linhas[7], linhas[8]
    assert total_op["custo"] == pytest.approx(31.5 + 1010.25 + 3.0)
    assert total_op["tma"] is None
    assert sem_der["custo"] == pytest.approx(120.5)
    assert total["custo"] == pytest.approx(total_op["custo"] + 120.5)
    assert total["clientes"] == 2
    assert total["custo_cpf"] == pytest.approx(total["custo"] / 2)
    assert total["prod_lead"] == pytest.approx(13000 / 40)


def test_custo_por_cpf_e_producao_por_lead_sem_divisao_por_zero(resultado):
    ana = next(l for l in ru.linhas_do_relatorio(resultado) if l["nome"] == "ANA ALVES")
    assert ana["custo_cpf"] == 0 and ana["prod_lead"] == 0


def test_operador_sem_atividade_nao_aparece(entrada, tmp_path):
    escrever_csv(entrada / "OPERADORES.csv", [
        ["", "? Nome", "Supervisor(es)", "Status"],
        ["", "ANA ALVES", "(1) SUP UM", "ATIVO"], ["", "BRUNO BRITO", "(1) SUP UM", "ATIVO"],
        ["", "CARLA CRUZ", "(1) CHEFE DOIS", "ATIVO"], ["", "ELIS PARADA", "(1) SUP UM", "ATIVO"],
    ])
    cadastro = tmp_path / "op.xlsx"
    entradas, _ = ru.ler_entradas(entrada)
    ru.atualizar_cadastro(entradas["operadores"][0][1], cadastro)
    resultado = ru.calcular(DIA, entradas, *ru.ler_cadastro(cadastro))
    assert "ELIS PARADA" not in [l["nome"] for l in ru.linhas_do_relatorio(resultado)]
    assert "ELIS PARADA" in [l["nome"] for l in ru.linhas_do_relatorio(resultado, mostrar_todos=True)]


def test_execucao_completa_gera_a_planilha(entrada, tmp_path, capsys):
    saida = tmp_path / "saida"
    cadastro = tmp_path / "config" / "operadores.xlsx"
    assert ru.main(["--data", "19/10/2026", "--sem-api", "--entrada", str(entrada),
                    "--saida", str(saida), "--cadastro", str(cadastro)]) == 0

    wb = openpyxl.load_workbook(saida / "relatorio_produtividade_2026-10-19.xlsx")
    ws = wb["Relatório"]
    assert ws["B2"].value == "RELATORIO DE PRODUTIVIDADE URA 19/10/2026"
    assert [ws.cell(3, c).value for c in range(2, 11)] == [
        "OPERADORES", "TMA", "TMO", "LEADS ATENDIDOS/URA", "CUSTO TOTAL", "VENDAS GERAL",
        "CLIENTE EFETIVADOS", "CUSTO POR CPF", "PRODUÇÃO/LEAD"]
    assert ws["B4"].value == "BRUNO BRITO"
    assert ws["C4"].number_format == "[h]:mm:ss"
    assert "Conferência" in wb.sheetnames

    saida_terminal = capsys.readouterr().out
    assert "NOMES SEM VÍNCULO" in saida_terminal and "KARLA CRUZ" in saida_terminal  # sem apelido aqui
    assert "Ex-colaboradores ignorados: DIEGO DIAS" in saida_terminal
    assert "Clientes efetivados: 2" in saida_terminal


def test_sem_cadastro_explica_o_que_fazer(tmp_path):
    with pytest.raises(FileNotFoundError, match="OPERADORES"):
        ru.ler_cadastro(tmp_path / "nao_existe.xlsx")
