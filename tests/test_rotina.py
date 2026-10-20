"""
Testes das funções da rotina diária: calendário de feriados, histórico e comparação
do relatório, metas (semáforo), arquivamento da entrada/ e o comando único benzaiten.py.

Todos os dados são INVENTADOS e criados em pastas temporárias.
"""

from datetime import date, time

import openpyxl
import pytest

import argus_api
import benzaiten
import calendario
import escala_bases
import relatorio_ura as ru

DIA = date(2026, 10, 19)  # segunda-feira


def escrever_csv(caminho, linhas):
    caminho.write_text("\r\n".join(";".join(f'"{c}"' for c in l) for l in linhas) + "\r\n", encoding="cp1252")


def montar_entrada(pasta, dia_texto="19/10/2026", leads_ana="30"):
    """Exportações mínimas de um dia: 2 operadores, tarifação e esteira."""
    pasta.mkdir(parents=True, exist_ok=True)
    escrever_csv(pasta / "OPERADORES.csv", [
        ["", "? Nome", "Supervisor(es)", "Status"],
        ["", "ANA ALVES", "(1) SUP UM", "ATIVO"], ["", "BRUNO BRITO", "(1) SUP UM", "ATIVO"]])
    escrever_csv(pasta / "tme_tma.csv", [
        ["Data", "Operador", "Supervisor", "TME", "TMA"],
        [dia_texto, "ANA ALVES", "SUP UM", "00:01:00", "00:02:00"],
        [dia_texto, "BRUNO BRITO", "SUP UM", "00:00:30", "00:00:40"]])
    escrever_csv(pasta / "leads.csv", [
        ["Data", "Operador", "Supervisor", "Leads. Atend."],
        [dia_texto, "ANA ALVES", "SUP UM", leads_ana], [dia_texto, "BRUNO BRITO", "SUP UM", "10"]])
    escrever_csv(pasta / "tarifacao.csv", [
        ["Descrição", "Valor R$", "Qtde. Lig", "Minutos"],
        ["<LIGAÇÕES SEM DERIVAÇÃO P/ OPERADOR>", "10,00", "1", "1"],
        ["ANA ALVES", "2,00", "1", "1"], ["BRUNO BRITO", "3,00", "1", "1"]])
    escrever_csv(pasta / "esteira.csv", [
        ["Data/Hora", "Etapa", "CPF", "Agente", "Valor Contrato"],
        [f"{dia_texto} 10:00:00", "Pago", "CLIENTE-000001", "BRUNO BRITO", "1000,00"]])


# --- Calendário ---------------------------------------------------------------

@pytest.mark.parametrize("ano, pascoa", [(2024, date(2024, 3, 31)), (2025, date(2025, 4, 20)),
                                         (2026, date(2026, 4, 5)), (2027, date(2027, 3, 28))])
def test_domingo_de_pascoa(ano, pascoa):
    assert calendario.domingo_de_pascoa(ano) == pascoa


def test_feriados_nacionais_incluem_sexta_santa_e_consciencia_negra():
    feriados = calendario.feriados_nacionais(2026)
    assert feriados[date(2026, 4, 3)] == "Sexta-feira Santa"
    assert date(2026, 11, 20) in feriados
    assert len(feriados) == 10


@pytest.mark.parametrize("hoje, esperado", [
    (date(2026, 10, 20), date(2026, 10, 19)),  # terça -> segunda
    (date(2026, 10, 19), date(2026, 10, 16)),  # segunda -> sexta
    (date(2026, 10, 13), date(2026, 10, 9)),   # terça depois do feriado de 12/10 (segunda) -> sexta
    (date(2026, 4, 6), date(2026, 4, 2)),      # segunda depois da Páscoa -> quinta (sexta é Santa)
])
def test_dia_util_anterior_pula_fim_de_semana_e_feriado(hoje, esperado):
    assert calendario.dia_util_anterior(hoje) == esperado


def test_feriado_local_da_planilha(tmp_path):
    arquivo = tmp_path / "feriados.xlsx"
    assert calendario.criar_modelo_feriados(arquivo) is True
    assert calendario.criar_modelo_feriados(arquivo) is False  # não sobrescreve
    wb = openpyxl.load_workbook(arquivo)
    wb.active.append(["16/10/2026", "Aniversário da cidade"])
    wb.active.append([date(2026, 10, 15), None])
    wb.save(arquivo)

    extras = calendario.ler_feriados(arquivo)
    assert extras == {date(2026, 10, 16): "Aniversário da cidade", date(2026, 10, 15): "Feriado"}
    # Segunda 19/10: sexta 16 e quinta 15 são feriados locais -> quarta 14
    assert calendario.dia_util_anterior(date(2026, 10, 19), extras) == date(2026, 10, 14)


# --- Histórico e comparação ------------------------------------------------------

def linha(tipo, nome, leads, custo, vendas):
    return {"tipo": tipo, "nome": nome, "tma": 60, "tmo": 30, "leads": leads,
            "custo": custo, "vendas": vendas, "clientes": 1}


def test_historico_grava_le_e_substitui_o_mesmo_dia(tmp_path):
    arquivo = tmp_path / "hist.csv"
    ru.salvar_no_historico(DIA, [linha("operador", "ANA", 5, 1.5, 1000.25), {"tipo": "separador", "nome": ""}], arquivo)
    ru.salvar_no_historico(DIA, [linha("operador", "ANA", 7, 2.0, 50.0)], arquivo)  # rodou de novo

    historico = ru.ler_historico(arquivo)
    assert len(historico) == 1
    assert historico[0] == {"data": DIA, "tipo": "operador", "nome": "ANA", "tma": 60, "tmo": 30,
                            "leads": 7, "custo": 2.0, "vendas": 50.0, "clientes": 1}
    assert "2,00" in arquivo.read_text(encoding="utf-8-sig")  # vírgula decimal para o Excel


def test_comparacao_usa_o_ultimo_relatorio_antes_do_dia(tmp_path):
    arquivo = tmp_path / "hist.csv"
    ru.salvar_no_historico(date(2026, 10, 15), [linha("operador", "ANA", 1, 1.0, 1.0)], arquivo)
    ru.salvar_no_historico(date(2026, 10, 16), [linha("operador", "ANA", 10, 5.0, 100.0)], arquivo)
    ru.salvar_no_historico(date(2026, 10, 20), [linha("operador", "ANA", 99, 99.0, 99.0)], arquivo)  # futuro

    linhas = [linha("operador", "Ána", 12, 4.0, 150.0), linha("operador", "NOVATA", 3, 1.0, 0.0)]
    anterior = ru.comparar_com_anterior(linhas, ru.ler_historico(arquivo), DIA)

    assert anterior == date(2026, 10, 16)
    assert (linhas[0]["d_leads"], linhas[0]["d_vendas"], linhas[0]["d_custo"]) == (2, 50.0, -1.0)
    assert "d_leads" not in linhas[1]  # não existia no relatório anterior


def test_sem_historico_nao_compara():
    assert ru.comparar_com_anterior([linha("operador", "ANA", 1, 1, 1)], [], DIA) is None


# --- Metas -------------------------------------------------------------------------

@pytest.mark.parametrize("valor, cor", [(40, "verde"), (30, "verde"), (20, "amarelo"), (10, "vermelho"), (0, "vermelho")])
def test_semaforo_maior_e_melhor(valor, cor):
    assert ru.cor_da_meta(valor, ("maior", 30, 10)) == cor


@pytest.mark.parametrize("valor, cor", [(50, "verde"), (60, "verde"), (90, "amarelo"), (120, "vermelho")])
def test_semaforo_menor_e_melhor(valor, cor):
    assert ru.cor_da_meta(valor, ("menor", 60, 120)) == cor


def test_semaforo_sem_regra_ou_sem_valor():
    assert ru.cor_da_meta(10, None) is None
    assert ru.cor_da_meta(None, ("maior", 1, 0)) is None
    assert ru.cor_da_meta(5, ("maior", 10, None)) == "amarelo"  # só o verde preenchido


def test_ler_metas_entende_tempo_e_valor_br(tmp_path):
    arquivo = tmp_path / "metas.xlsx"
    assert ru.criar_modelo_metas(arquivo) is True
    wb = openpyxl.load_workbook(arquivo)
    ws = wb["Metas"]
    valores = {"TMA": ("00:01:00", time(0, 2, 0)), "TMO": (45, 1 / 24 / 60 * 1.5),  # 1,5 min como fração do dia
               "VENDAS GERAL": ("15.000,50", "5000"), "CUSTO POR CPF": (None, None)}
    for celulas in ws.iter_rows(min_row=2, max_col=4):
        if celulas[0].value in valores:
            celulas[2].value, celulas[3].value = valores[celulas[0].value]
    wb.save(arquivo)

    metas = ru.ler_metas(arquivo)
    assert metas["tma"] == ("menor", 60, 120)
    assert metas["tmo"] == ("menor", 45, 90)
    assert metas["vendas"] == ("maior", 15000.5, 5000.0)
    assert "custo_cpf" not in metas  # em branco = não pinta


# --- Arquivamento -----------------------------------------------------------------

def test_arquivar_move_sem_alterar_e_sem_sobrescrever(tmp_path):
    entrada = tmp_path / "entrada"
    entrada.mkdir()
    (entrada / "tarifa.csv").write_text("conteúdo original", encoding="utf-8")
    destino = entrada / "processados" / "2026-10-19"
    destino.mkdir(parents=True)
    (destino / "tarifa.csv").write_text("arquivado antes", encoding="utf-8")

    pasta = ru.arquivar([entrada / "tarifa.csv", entrada / "tarifa.csv"], entrada, DIA)

    assert pasta == destino
    assert not (entrada / "tarifa.csv").exists()
    assert (destino / "tarifa.csv").read_text(encoding="utf-8") == "arquivado antes"
    assert (destino / "tarifa_2.csv").read_text(encoding="utf-8") == "conteúdo original"


# --- Relatório com histórico, comparação, metas e arquivamento ---------------------

@pytest.fixture
def projeto(tmp_path, monkeypatch):
    """Aponta todos os scripts para uma pasta temporária."""
    monkeypatch.setattr(ru, "PASTA_ENTRADA", tmp_path / "entrada")
    monkeypatch.setattr(ru, "PASTA_SAIDA", tmp_path / "saida")
    monkeypatch.setattr(ru, "ARQUIVO_CADASTRO", tmp_path / "config" / "operadores.xlsx")
    monkeypatch.setattr(ru, "ARQUIVO_METAS", tmp_path / "config" / "metas.xlsx")
    monkeypatch.setattr(ru, "ARQUIVO_HISTORICO", tmp_path / "historico" / "relatorio_ura.csv")
    monkeypatch.setattr(calendario, "ARQUIVO_FERIADOS", tmp_path / "config" / "feriados.xlsx")
    monkeypatch.setattr(escala_bases, "ARQUIVO_CADASTRO", tmp_path / "config" / "bases.xlsx")
    monkeypatch.setattr(escala_bases, "ARQUIVO_HISTORICO", tmp_path / "historico" / "historico_escalas.csv")
    monkeypatch.setattr(escala_bases, "PASTA_SAIDA", tmp_path / "saida")
    monkeypatch.setattr(argus_api, "ARQUIVO_TOKEN", tmp_path / "config" / "sem_token.txt")
    monkeypatch.delenv("ARGUS_TOKEN", raising=False)
    return tmp_path


def criar_bases(pasta):
    import pandas as pd

    (pasta / "config").mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(pasta / "config" / "bases.xlsx") as escritor:
        pd.DataFrame({"Base": ["BASE A", "BASE B"]}).to_excel(escritor, sheet_name="Bases", index=False)
        pd.DataFrame({"Pessoa": ["ANA", "BRUNO"]}).to_excel(escritor, sheet_name="Pessoas", index=False)


def test_relatorio_compara_pinta_e_arquiva(projeto):
    # Dia 1: sem comparação
    montar_entrada(projeto / "entrada", "16/10/2026", leads_ana="20")
    escrever_csv(projeto / "entrada" / "tme_tma_velho.csv", [  # de outro dia: não deve ser arquivado
        ["Data", "Operador", "Supervisor", "TME", "TMA"], ["01/10/2026", "ANA ALVES", "SUP UM", "00:00:01", "00:00:01"]])
    assert ru.main(["--data", "16/10/2026", "--sem-api", "--arquivar"]) == 0
    entrada = projeto / "entrada"
    assert sorted(p.name for p in entrada.iterdir() if p.is_file()) == ["tme_tma_velho.csv"]
    assert len(list((entrada / "processados" / "2026-10-16").iterdir())) == 5

    # Metas: leads verde a partir de 30, vermelho até 10
    wb = openpyxl.load_workbook(ru.ARQUIVO_METAS)
    for celulas in wb["Metas"].iter_rows(min_row=2, max_col=4):
        if celulas[0].value == "LEADS ATENDIDOS/URA":
            celulas[2].value, celulas[3].value = 30, 10
    wb.save(ru.ARQUIVO_METAS)

    # Dia 2: ANA atendeu 30 leads (eram 20)
    montar_entrada(entrada, "19/10/2026", leads_ana="30")
    assert ru.main(["--data", "19/10/2026", "--sem-api"]) == 0

    ws = openpyxl.load_workbook(projeto / "saida" / "relatorio_produtividade_2026-10-19.xlsx")["Relatório"]
    titulos = [ws.cell(3, c).value for c in range(2, 14)]
    assert titulos[-3:] == ["LEADS vs 16/10", "VENDAS vs 16/10", "CUSTO vs 16/10"]
    linhas = {ws.cell(r, 2).value: r for r in range(4, ws.max_row + 1) if ws.cell(r, 2).value}
    ana, bruno = linhas["ANA ALVES"], linhas["BRUNO BRITO"]
    assert ws.cell(ana, 11).value == 10                       # +10 leads
    assert ws.cell(ana, 11).font.color.rgb.endswith("006100")  # verde: subiu e é bom
    assert ws.cell(ana, 5).fill.fgColor.rgb.endswith("C6EFCE")  # 30 leads: meta verde
    assert ws.cell(bruno, 5).fill.fgColor.rgb.endswith("FFC7CE")  # 10 leads: meta vermelha
    assert len({r["data"] for r in ru.ler_historico(ru.ARQUIVO_HISTORICO)}) == 2


# --- Comando único (benzaiten.py) ---------------------------------------------------

def test_benzaiten_faz_escala_e_relatorio(projeto, capsys):
    criar_bases(projeto)
    montar_entrada(projeto / "entrada", "19/10/2026")

    assert benzaiten.main(["--hoje", "20/10/2026"]) == 0  # terça: relatório de segunda 19/10

    saida = capsys.readouterr().out
    assert "relatório de 19/10/2026" in saida
    assert "[ OK ] Escala de bases: sorteada" in saida
    assert "[ OK ] Relatório da URA: gerado e arquivos arquivados" in saida
    assert (projeto / "saida" / "escala_2026-10-20.xlsx").exists()
    assert (projeto / "saida" / "relatorio_produtividade_2026-10-19.xlsx").exists()
    assert (projeto / "config" / "feriados.xlsx").exists() and (projeto / "config" / "metas.xlsx").exists()
    assert not list((projeto / "entrada").glob("*.csv"))  # tudo arquivado


def test_benzaiten_nao_sorteia_duas_vezes_nem_em_fim_de_semana(projeto, capsys):
    criar_bases(projeto)
    benzaiten.main(["--hoje", "20/10/2026", "--sem-relatorio"])
    benzaiten.main(["--hoje", "20/10/2026", "--sem-relatorio"])
    benzaiten.main(["--hoje", "24/10/2026", "--sem-relatorio"])  # sábado
    saida = capsys.readouterr().out
    assert "Escala de bases: já estava sorteada" in saida
    assert "Escala de bases: pulada (não é dia útil)" in saida
    assert len(escala_bases.ler_historico(escala_bases.ARQUIVO_HISTORICO)) == 2  # só o dia 20


def test_benzaiten_avisa_o_que_falta_e_continua_se_uma_etapa_falhar(projeto, capsys):
    # Sem config/bases.xlsx: a escala falha, mas o relatório roda
    montar_entrada(projeto / "entrada", "19/10/2026")
    (projeto / "entrada" / "esteira.csv").unlink()

    assert benzaiten.main(["--hoje", "20/10/2026", "--nao-arquivar"]) == 1

    saida = capsys.readouterr().out
    assert "PENDENTE: exportar a esteira do Vanguard" in saida
    assert "PENDENTE: exportar as tarifações VENDITORE e 4NET" in saida  # só 1 arquivo
    assert "[ERRO] Escala de bases" in saida
    assert "[ OK ] Relatório da URA: gerado" in saida
    assert (projeto / "entrada" / "tarifacao.csv").exists()  # --nao-arquivar
