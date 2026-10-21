"""
Testes do ranking de contratos pagos (ranking_vendas.py) e das regras da esteira (esteira.py).

Todos os nomes, valores e metas são INVENTADOS e criados dentro do teste.
"""

from datetime import date

import openpyxl
import pytest

import ranking_vendas as rv
import relatorio_ura as ru
from esteira import eh_cartao, nome_do_agente, remover_propostas_repetidas

DIA = date(2026, 10, 20)       # terça; outubro/2026 tem 21 dias úteis (12/10 é feriado)
FATOR = 13 / 21                # 13 dias úteis de 01/10 a 20/10

METAS = {
    "meta_vendedor": {"CLT": 1000, "ESTAGIO": 800},
    "bonus_vendedor": {"CLT": 600, "ESTAGIO": 500},
    "meta_supervisor": {"CLT": 5000, "ESTAGIO": 4000},
    "bonus_supervisor": {"CLT": 3000, "ESTAGIO": 2500},
    "meta_geral": {"CLT": 9000},
    "meta_cartao": {"CLT": 5},
    "emoji": {"CLT": "*"},
}
CABECALHO = ["Proposta", "Data/Hora", "Data Pagamento", "Etapa", "CPF", "Produto", "Agente", "Equipe", "Valor Contrato"]


def escrever_csv(caminho, linhas, codificacao="utf-8-sig"):
    caminho.write_text("\r\n".join(";".join(f'"{c}"' for c in l) for l in linhas) + "\r\n", encoding=codificacao)


def proposta(numero, etapa, agente, valor, equipe="SUP UM", produto="PORT COM REFIN", pago_em="15/10/2026"):
    return [str(numero), "01/10/2026 09:00:00", pago_em if etapa == "Pago" else "", etapa,
            f"CLIENTE-{numero:06d}", produto, agente, equipe, valor]


@pytest.fixture
def projeto(tmp_path):
    """Pasta com cadastro, config do ranking e as duas esteiras (andamento e pagamento)."""
    entrada = tmp_path / "entrada"
    entrada.mkdir()
    escrever_csv(entrada / "OPERADORES.csv", [
        ["", "? Nome", "Supervisor(es)", "Status"],
        ["", "ANA ALVES", "(1) SUP UM", "ATIVO"],
        ["", "BRUNO BRITO", "(1) SUP UM", "ATIVO"],
        ["", "CARLA CRUZ", "(1) SUP DOIS", "ATIVO"],
        ["", "DIEGO DIAS", "(1) SUP DOIS", "INATIVO"],
    ], "cp1252")
    escrever_csv(entrada / "esteira_pagamento.csv", [CABECALHO,
        proposta(1, "Pago", "ANA ALVES - COMERCIAL", "600,00"),
        proposta(2, "Pago", "ANA ALVES", "100,00", produto="MARGEM LIVRE"),        # margem: só conta cartão
        proposta(3, "Pago", "BRUNO BRITO", "1.000,00"),
        proposta(4, "Pago", "BRUNO BRITO", "999,00", pago_em="30/09/2026"),      # pago em setembro: fora
        proposta(5, "Pago", "CARLA CRUZ", "50,00", equipe="SUP DOIS"),
        proposta(6, "Pago", "DIEGO DIAS", "777,00", equipe="SUP DOIS"),          # inativo: fora do ranking
    ])
    escrever_csv(entrada / "esteira_andamento.csv", [CABECALHO,
        proposta(1, "Andamento", "ANA ALVES", "600,00"),                          # repetida: já está paga
        proposta(7, "Andamento", "ANA ALVES", "200,00"),
        proposta(8, "Andamento", "CARLA CRUZ", "300,00", equipe="SUP DOIS"),
        proposta(9, "Pendente", "CARLA CRUZ", "5000,00", equipe="SUP DOIS"),     # Pendente não conta
    ])
    config = tmp_path / "config"
    config.mkdir()
    rv.salvar_config(config / "ranking.xlsx", METAS, ["BRUNO BRITO"], ["SUP UM"], {"SUP DOIS": "Equipe da Carla"})
    entradas, _ = ru.ler_entradas(entrada)
    ru.atualizar_cadastro(entradas["operadores"][0][1], config / "operadores.xlsx")
    return tmp_path


def calcular(projeto):
    entradas, _ = ru.ler_entradas(projeto / "entrada")
    ativos, inativos, apelidos = ru.ler_cadastro(projeto / "config" / "operadores.xlsx")
    config = rv.ler_config(projeto / "config" / "ranking.xlsx")
    return rv.calcular(DIA, entradas, ativos, inativos, apelidos, config, {})


def por_nome(linhas):
    return {l.nome: l for l in linhas}


# --- Regras da esteira ---------------------------------------------------------

def test_proposta_repetida_fica_uma_vez_e_prefere_a_paga():
    linhas = [{"PROPOSTA": "1", "ETAPA": "Andamento"}, {"PROPOSTA": "1", "ETAPA": "Pago"},
              {"PROPOSTA": "", "NUMERO-PROPOSTA": "", "ETAPA": "Pago"}, {"PROPOSTA": "2", "ETAPA": "Pago"}]
    unicas, repetidas = remover_propostas_repetidas(linhas)
    assert repetidas == 1
    assert [l["ETAPA"] for l in unicas] == ["Pago", "Pago", "Pago"]  # sem número: mantida


def test_nome_do_agente_e_produto_cartao():
    assert nome_do_agente("  ANA  ALVES - COMERCIAL ") == "ANA ALVES"
    assert eh_cartao({"PRODUTO": "Cartão Benefício"}) and not eh_cartao({"PRODUTO": "PORT COM REFIN"})


# --- Calendário do mês -----------------------------------------------------------

def test_dias_uteis_de_outubro_de_2026():
    assert rv.dias_uteis(date(2026, 10, 1), date(2026, 10, 31), {}) == 21
    assert rv.dias_uteis(date(2026, 10, 1), DIA, {}) == 13
    assert rv.fim_do_mes(date(2026, 2, 10)) == date(2026, 2, 28)


# --- Bônus -------------------------------------------------------------------------

@pytest.mark.parametrize("atingimento, taxa, ultima", [(0.5, 0.0015, False), (1.1, 0.0025, False), (1.3, 0.003, True)])
def test_faixas_do_bonus_clt(atingimento, taxa, ultima):
    bonus, na_ultima = rv.bonus_da_faixa(atingimento, 1000, rv.FAIXAS_VENDEDOR["CLT"])
    assert bonus == pytest.approx(1000 * atingimento * taxa)
    assert na_ultima is ultima


# --- Ranking ---------------------------------------------------------------------

def test_portabilidade_cartao_e_andamento(projeto):
    vend = por_nome(calcular(projeto).vendedores)
    ana = vend["ANA ALVES"]
    assert ana.portabilidade == 600         # "ANA ALVES - COMERCIAL" vira ANA ALVES
    assert (ana.qtd_cartao, ana.valor_cartao) == (1, 100)
    assert ana.andamento == 200             # a proposta 1 já paga não entra em andamento
    assert ana.projecao == 600 + 200 * 0.5
    assert vend["BRUNO BRITO"].portabilidade == 1000  # o pago de setembro ficou fora
    assert vend["CARLA CRUZ"].andamento == 300         # Pendente não conta


def test_metas_clt_e_estagio_projetando_potencial_e_elegivel(projeto):
    vend = por_nome(calcular(projeto).vendedores)
    ana, bruno = vend["ANA ALVES"], vend["BRUNO BRITO"]

    # ANA é estágio (meta 800): 600 / (800 x 13/21)
    assert ana.categoria == "ESTAGIO"
    assert ana.projetando == pytest.approx(600 / (800 * FATOR))
    assert ana.potencial == pytest.approx((800 * ana.projetando * 0.0025 + 100 * 0.01) * 1.2)
    assert ana.elegivel is True             # 600 >= 500

    # BRUNO é CLT (meta 1000): 1000 / (1000 x 13/21) = 1,615 -> última faixa, com extra de R$ 1.000
    assert bruno.categoria == "CLT"
    assert bruno.projetando == pytest.approx(21 / 13)
    assert bruno.potencial == pytest.approx((1000 * bruno.projetando * 0.003 + 1000) * 1.2)


def test_ordem_pelo_projetando(projeto):
    nomes = [l.nome for l in calcular(projeto).vendedores]
    assert nomes == ["BRUNO BRITO", "ANA ALVES", "CARLA CRUZ"]
    assert [l.posicao for l in calcular(projeto).vendedores] == [1, 2, 3]


def test_empate_divide_a_posicao():
    linhas = rv.classificar([rv.Linha("A", projetando=0), rv.Linha("B", projetando=0.5), rv.Linha("C", projetando=0)])
    assert [(l.nome, l.posicao) for l in linhas] == [("B", 1), ("A", 2), ("C", 2)]


def test_inativo_e_etapas_fora_viram_aviso(projeto):
    ranking = calcular(projeto)
    assert "DIEGO DIAS" not in por_nome(ranking.vendedores)
    assert ranking.ignorados == {"DIEGO DIAS": 777.0}
    assert ranking.contagem["Propostas repetidas entre as esteiras (contadas uma vez)"] == 1
    assert ranking.contagem["Pagas fora do mês (ignoradas)"] == 1
    assert any("Pendente (1)" in a for a in ranking.avisos)


def test_supervisores_nomes_metas_e_total_geral(projeto):
    ranking = calcular(projeto)
    sup = por_nome(ranking.supervisores)
    # SUP DOIS inclui o DIEGO (inativo): na macro, a equipe soma todas as propostas
    assert sup["Equipe da Carla"].portabilidade == 50 + 777
    assert sup["Equipe da Carla"].categoria == "ESTAGIO"
    assert sup["SUP UM"].categoria == "CLT"
    assert sup["SUP UM"].portabilidade == 1600
    total = ranking.total
    assert total["portabilidade"] == 1600 + 827
    assert total["projetando"] == pytest.approx(2427 / (9000 * FATOR))
    assert total["elegiveis"] == sum(l.elegivel for l in ranking.vendedores + ranking.supervisores)


def test_esteira_com_coluna_a_mais_no_meio_da_o_mesmo_resultado(projeto):
    """O erro da planilha antiga: a esteira nova ganhou uma coluna e a macro lia a coluna errada."""
    antes = {l.nome: l.portabilidade for l in calcular(projeto).vendedores}
    for nome in ("esteira_pagamento.csv", "esteira_andamento.csv"):
        arquivo = projeto / "entrada" / nome
        linhas = [l.split(";") for l in arquivo.read_text(encoding="utf-8-sig").splitlines()]
        novas = [l[:5] + (['"Codigo"'] if i == 0 else ['"X"']) + l[5:] for i, l in enumerate(linhas)]
        arquivo.write_text("\r\n".join(";".join(l) for l in novas) + "\r\n", encoding="utf-8-sig")
    assert {l.nome: l.portabilidade for l in calcular(projeto).vendedores} == antes


# --- Feedback ------------------------------------------------------------------------

def test_feedback_respeitoso_e_trata_pelo_primeiro_nome():
    assert rv.como_chamar("JOANA DE SOUZA PRADO") == "Joana"
    assert rv.como_chamar("SEM SUPERVISOR") == "Sem Supervisor"
    assert rv.como_chamar("FULANO - URA") == "Fulano - Ura"
    textos = [
        rv.gerar_feedback(rv.Linha("ANA X", projetando=1.5, portabilidade=1500), 1000),
        rv.gerar_feedback(rv.Linha("ANA X", projetando=1.0, portabilidade=1000), 1000),
        rv.gerar_feedback(rv.Linha("ANA X", projetando=0.8, portabilidade=800), 1000),
        rv.gerar_feedback(rv.Linha("ANA X", projetando=0.3, portabilidade=300, andamento=500), 1000),
        rv.gerar_feedback(rv.Linha("ANA X", projetando=0.3, portabilidade=300), 1000),
        rv.gerar_feedback(rv.Linha("ANA X", andamento=500), 1000),
        rv.gerar_feedback(rv.Linha("ANA X"), 1000),
    ]
    assert all(t.startswith("Ana, ") for t in textos)
    assert "faltam R$ 200" in textos[2]
    assert "R$ 0" not in " ".join(textos)


# --- Configuração e execução completa --------------------------------------------------

def test_config_inexistente_ou_incompleta_explica(tmp_path):
    with pytest.raises(FileNotFoundError, match="--importar-xlsm"):
        rv.ler_config(tmp_path / "nao_existe.xlsx")
    metas = dict(METAS, meta_geral={"CLT": None})
    rv.salvar_config(tmp_path / "r.xlsx", metas, [], [], {})
    with pytest.raises(ValueError, match="META GERAL"):
        rv.ler_config(tmp_path / "r.xlsx")


def test_importar_planilha_antiga(tmp_path):
    wb = openpyxl.Workbook()
    guia = wb.active
    guia.title = "GUIA"
    for celula, valor in {"B7": 5, "E7": 5, "B8": 1000, "E8": 800, "B10": 600, "E10": 500, "B11": 3000,
                          "E11": 2500, "B14": 5000, "E14": 4000, "B17": 9000, "B19": "*"}.items():
        guia[celula] = valor
    aux = wb.create_sheet("AUXILIAR")
    aux["A1"], aux["B1"], aux["C1"] = "SUP UM", "SUPERVISOR UM", "Tabela de Supervisor CLT"
    aux["C2"] = "SUP DOIS"
    aux["A7"] = "ESSE TRECHO AQUI EM CIMA ^"
    aux["A13"], aux["A14"], aux["A15"] = "FUNCIONÁRIOS CLT", "BRUNO BRITO", "CARLA  CRUZ"
    antiga = tmp_path / "antiga.xlsx"
    wb.save(antiga)

    destino = tmp_path / "ranking.xlsx"
    assert rv.importar_xlsm(antiga, destino) is True
    assert rv.importar_xlsm(antiga, destino) is False  # não sobrescreve
    config = rv.ler_config(destino)
    assert config.meta("meta_vendedor", "ESTAGIO") == 800
    assert config.vendedores_clt == {"BRUNO BRITO", "CARLA CRUZ"}
    assert config.supervisores_clt == {"SUP DOIS"}
    assert config.nomes_equipes == {"SUP UM": "SUPERVISOR UM"}


def test_execucao_completa_gera_a_planilha(projeto, capsys):
    saida = projeto / "saida"
    usados = []
    assert rv.main(["--data", "20/10/2026", "--hoje", "21/10/2026", "--entrada", str(projeto / "entrada"),
                    "--saida", str(saida), "--cadastro", str(projeto / "config" / "operadores.xlsx"),
                    "--config", str(projeto / "config" / "ranking.xlsx")], usados) == 0

    ws = openpyxl.load_workbook(saida / "ranking_vendas_2026-10-20.xlsx")["Ranking"]
    assert ws["B2"].value == "RANKING GERAL DE CONTRATOS PAGOS - VENDEDOR - DE 01/10/2026 A 20/10/2026"
    assert ws["I3"].value == "PORT. PROJEÇÃO - 28/10/2026"
    assert ws["C4"].value == "BRUNO BRITO   *"   # >= 133%: ganha o emoji
    assert ws["J4"].value == "SIM"
    assert sorted(p.name for p in usados) == ["esteira_andamento.csv", "esteira_pagamento.csv"]
    assert "Top 3 vendedores" in capsys.readouterr().out
