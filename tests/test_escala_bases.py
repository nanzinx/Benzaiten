"""
Testes do sorteio de bases (escala_bases.py).

Todos os dados aqui são FICTÍCIOS e criados dentro do próprio teste,
numa pasta temporária. Nada é lido de config/ ou historico/.
"""

from collections import Counter
from datetime import date

import openpyxl
import pandas as pd
import pytest

import escala_bases as eb

BASES = ["BASE A", "BASE B", "BASE C", "BASE D"]
PESSOAS = ["ANA", "BRUNO", "CARLA", "DIEGO"]


def registro(d, base, pessoa):
    return {"data": d, "base": base, "responsavel": pessoa}


def historico_de(*dias):
    """Monta um histórico a partir de pares (data, {base: pessoa})."""
    regs = []
    for d, escolha in dias:
        regs += [registro(d, b, p) for b, p in escolha.items()]
    return regs


# Dois sorteios anteriores (sexta 16/10 e segunda 19/10)
HISTORICO = historico_de(
    (date(2026, 10, 16), {"BASE A": "ANA", "BASE B": "BRUNO", "BASE C": "CARLA", "BASE D": "DIEGO"}),
    (date(2026, 10, 19), {"BASE A": "BRUNO", "BASE B": "ANA", "BASE C": "DIEGO", "BASE D": "CARLA"}),
)


# --- Cadastro ---------------------------------------------------------------

def test_ler_cadastro_ignora_vazios_espacos_e_repetidos(tmp_path):
    arquivo = tmp_path / "bases.xlsx"
    with pd.ExcelWriter(arquivo) as escritor:
        pd.DataFrame({"Base": ["BASE A", None, "  BASE B ", "base a"]}).to_excel(escritor, sheet_name="Bases", index=False)
        pd.DataFrame({"Pessoa": ["ANA", "", "BRUNO"]}).to_excel(escritor, sheet_name="Pessoas", index=False)

    bases, pessoas = eb.ler_cadastro(arquivo)

    assert bases == ["BASE A", "BASE B"]
    assert pessoas == ["ANA", "BRUNO"]


def test_ler_cadastro_inexistente_explica_o_que_fazer(tmp_path):
    with pytest.raises(FileNotFoundError, match="--importar-xlsm"):
        eb.ler_cadastro(tmp_path / "nao_existe.xlsx")


# --- Regras do rodízio ------------------------------------------------------

NOTAS = ["INSS - LOTE - 1 - NOTA 0", "INSS - LOTE - 1 - NOTA 1", "INSS - LOTE - 1 - NOTA 2", "INSS - LOTE - 1 - NOTA 3"]

# O caso que motivou a mudança: no sorteio aleatório, quem tinha a NOTA 0 no dia 19
# pegou a NOTA 0 de novo no dia 21 (a planilha só olhava o dia anterior).
CASO_REAL = historico_de(
    (date(2026, 10, 16), dict(zip(NOTAS, ["BRUNO", "DIEGO", "CARLA", "ANA"]))),
    (date(2026, 10, 19), dict(zip(NOTAS, ["ANA", "BRUNO", "DIEGO", "CARLA"]))),
    (date(2026, 10, 20), dict(zip(NOTAS, ["DIEGO", "CARLA", "ANA", "BRUNO"]))),
    (date(2026, 10, 21), dict(zip(NOTAS, ["ANA", "BRUNO", "CARLA", "DIEGO"]))),
)


def test_cada_pessoa_anda_uma_base_para_frente():
    # 21/10: ANA na NOTA 0, BRUNO na 1, CARLA na 2, DIEGO na 3
    escolha, avisos = eb.rodizio(NOTAS, PESSOAS, CASO_REAL, date(2026, 10, 22))

    assert list(escolha) == NOTAS                        # todas as bases, na ordem do cadastro
    assert escolha == dict(zip(NOTAS, ["DIEGO", "ANA", "BRUNO", "CARLA"]))
    assert avisos == []


def test_ninguem_repete_base_antes_de_passar_por_todas():
    historico, dia = list(CASO_REAL), date(2026, 10, 22)
    vistas = {p: [] for p in PESSOAS}
    for _ in range(8):                                   # dois ciclos completos
        escolha, _ = eb.rodizio(NOTAS, PESSOAS, historico, dia)
        historico += [registro(dia, b, p) for b, p in escolha.items()]
        for base, pessoa in escolha.items():
            vistas[pessoa].append(base)
        dia = date.fromordinal(dia.toordinal() + 1)
    for pessoa, sequencia in vistas.items():
        assert sorted(sequencia[:4]) == NOTAS, pessoa    # cada ciclo passa por todas as bases
        assert sequencia[4:] == sequencia[:4], pessoa    # e o ciclo seguinte repete a mesma ordem


def test_rodizio_conta_sorteios_e_nao_dias_do_calendario():
    # De sexta (16/10) para segunda (19/10) é só um passo
    sexta = historico_de((date(2026, 10, 16), dict(zip(BASES, PESSOAS))))
    escolha, _ = eb.rodizio(BASES, PESSOAS, sexta, date(2026, 10, 19))
    assert escolha == dict(zip(BASES, ["DIEGO", "ANA", "BRUNO", "CARLA"]))


def test_sortear_de_novo_o_mesmo_dia_da_o_mesmo_resultado():
    # O dia do sorteio e os dias futuros não contam como ponto de partida
    primeira, _ = eb.rodizio(NOTAS, PESSOAS, CASO_REAL[:12], date(2026, 10, 21))
    segunda, _ = eb.rodizio(NOTAS, PESSOAS, CASO_REAL[:12], date(2026, 10, 21))
    assert primeira == segunda == dict(zip(NOTAS, ["BRUNO", "DIEGO", "CARLA", "ANA"]))


def test_sem_historico_comeca_pela_ordem_do_cadastro():
    escolha, _ = eb.rodizio(BASES, PESSOAS, [], date(2026, 10, 20))
    assert escolha == dict(zip(BASES, PESSOAS))


def test_comparacao_com_historico_ignora_acentos_e_maiusculas():
    hist = historico_de((date(2026, 10, 19), {"base a": "Ána", "base b": "bruno", "base c": "CARLA", "base d": "diego"}))
    escolha, _ = eb.rodizio(BASES, PESSOAS, hist, date(2026, 10, 20))
    assert escolha == dict(zip(BASES, ["DIEGO", "ANA", "BRUNO", "CARLA"]))


def test_ausente_retoma_sua_posicao_quando_volta():
    # 22/10: ANA falta. A base dela (NOTA 1) vai para alguém; no dia 23 ela volta
    # exatamente onde estaria (NOTA 2), como se não tivesse faltado.
    escolha, avisos = eb.rodizio(NOTAS, PESSOAS, CASO_REAL, date(2026, 10, 22), ausentes=["ana"])
    assert "ANA" not in escolha.values()
    assert escolha[NOTAS[1]] in {"BRUNO", "CARLA", "DIEGO"}
    assert escolha[NOTAS[1]] != "BRUNO"                  # BRUNO teve a NOTA 1 no dia 21
    assert "ANA está ausente" in avisos[0]

    historico = CASO_REAL + [registro(date(2026, 10, 22), b, p) for b, p in escolha.items()]
    volta, _ = eb.rodizio(NOTAS, PESSOAS, historico, date(2026, 10, 23))
    assert volta == dict(zip(NOTAS, ["CARLA", "DIEGO", "ANA", "BRUNO"]))


def test_base_do_ausente_vai_para_quem_esta_com_menos_bases():
    escolha, _ = eb.rodizio(NOTAS, PESSOAS, CASO_REAL, date(2026, 10, 22), ausentes=["ANA", "BRUNO"])
    assert sorted(Counter(escolha.values()).values()) == [2, 2]


def test_pessoa_nova_entra_no_fim_da_roda():
    escolha, _ = eb.rodizio(NOTAS, PESSOAS + ["ELISA"], CASO_REAL, date(2026, 10, 22))
    # Roda: ANA, BRUNO, CARLA, DIEGO, ELISA. Quem vem antes da NOTA 0 agora é a ELISA
    assert escolha == dict(zip(NOTAS, ["ELISA", "ANA", "BRUNO", "CARLA"]))


def test_menos_pessoas_que_bases_distribui_sem_diferenca_maior_que_1():
    escolha, _ = eb.rodizio(BASES, ["ANA", "BRUNO", "CARLA"], [], date(2026, 10, 20))

    quantidade = Counter(escolha.values())
    assert len(escolha) == 4                      # nenhuma base sem responsável
    assert set(quantidade) == {"ANA", "BRUNO", "CARLA"}
    assert max(quantidade.values()) - min(quantidade.values()) <= 1


def test_com_2_pessoas_para_4_bases_alterna_a_cada_dia():
    hist = historico_de((date(2026, 10, 19), {"BASE A": "ANA", "BASE B": "BRUNO", "BASE C": "ANA", "BASE D": "BRUNO"}))
    escolha, _ = eb.rodizio(BASES, ["ANA", "BRUNO"], hist, date(2026, 10, 20))
    assert escolha == {"BASE A": "BRUNO", "BASE B": "ANA", "BASE C": "BRUNO", "BASE D": "ANA"}


def test_mais_pessoas_que_bases_quem_fica_sem_base_tambem_roda():
    pessoas = PESSOAS + ["ELISA"]
    historico, dia, sem_base = [], date(2026, 10, 20), []
    for _ in range(5):
        escolha, _ = eb.rodizio(BASES, pessoas, historico, dia)
        assert len(set(escolha.values())) == 4    # ninguém com duas bases
        sem_base += list(set(pessoas) - set(escolha.values()))
        historico += [registro(dia, b, p) for b, p in escolha.items()]
        dia = date.fromordinal(dia.toordinal() + 1)
    assert sorted(sem_base) == sorted(pessoas)    # cada um fica de fora uma vez a cada 5 dias


def test_sem_bases_ou_sem_pessoas_da_erro():
    with pytest.raises(ValueError):
        eb.rodizio([], PESSOAS, [], date(2026, 10, 20))
    with pytest.raises(ValueError):
        eb.rodizio(BASES, PESSOAS, [], date(2026, 10, 20), ausentes=PESSOAS)


# --- Ausências --------------------------------------------------------------

def test_ausentes_saem_do_sorteio_sem_diferenciar_maiusculas():
    assert eb.aplicar_ausencias(PESSOAS, ["carla", " Diego "]) == ["ANA", "BRUNO"]


def test_ausente_com_nome_errado_da_erro():
    with pytest.raises(ValueError, match="CARLLA"):
        eb.aplicar_ausencias(PESSOAS, ["CARLLA"])


# --- Histórico e execução completa -------------------------------------------

def test_historico_grava_e_le_igual(tmp_path):
    arquivo = tmp_path / "historico.csv"
    eb.salvar_historico(HISTORICO, arquivo)
    assert eb.ler_historico(arquivo) == HISTORICO


@pytest.fixture
def projeto_ficticio(tmp_path, monkeypatch):
    """Aponta o script para uma pasta temporária com um cadastro fictício."""
    cadastro = tmp_path / "config" / "bases.xlsx"
    cadastro.parent.mkdir()
    with pd.ExcelWriter(cadastro) as escritor:
        pd.DataFrame({"Base": BASES}).to_excel(escritor, sheet_name="Bases", index=False)
        pd.DataFrame({"Pessoa": PESSOAS}).to_excel(escritor, sheet_name="Pessoas", index=False)
    monkeypatch.setattr(eb, "ARQUIVO_CADASTRO", cadastro)
    monkeypatch.setattr(eb, "ARQUIVO_HISTORICO", tmp_path / "historico" / "historico_escalas.csv")
    monkeypatch.setattr(eb, "PASTA_SAIDA", tmp_path / "saida")
    return tmp_path


def test_execucao_completa_gera_planilha_e_historico(projeto_ficticio):
    assert eb.main(["--data", "21/10/2026", "--ausentes", "ana"]) == 0

    planilha = projeto_ficticio / "saida" / "escala_2026-10-21.xlsx"
    escala = openpyxl.load_workbook(planilha)["Escala"]
    assert escala["A1"].value == "ESCALA DE BASES — ATIVO 2"
    assert escala["A2"].value == "21/10/2026 (quarta-feira)"
    responsaveis = [escala.cell(row=r, column=2).value for r in range(5, 9)]
    assert "ANA" not in responsaveis
    assert any("Ausentes: ANA" == str(c.value) for c in escala["A"])

    historico = eb.ler_historico(eb.ARQUIVO_HISTORICO)
    assert len(historico) == 4
    assert eb.ler_cadastro(eb.ARQUIVO_CADASTRO) == (BASES, PESSOAS)  # cadastro não foi alterado


def test_data_ja_sorteada_nao_e_substituida_sem_confirmar(projeto_ficticio, monkeypatch):
    eb.main(["--data", "21/10/2026"])
    antes = eb.ler_historico(eb.ARQUIVO_HISTORICO)

    monkeypatch.setattr("builtins.input", lambda _: "n")
    assert eb.main(["--data", "21/10/2026"]) == 1
    assert eb.ler_historico(eb.ARQUIVO_HISTORICO) == antes


def test_substituir_troca_o_dia_sem_duplicar(projeto_ficticio):
    eb.main(["--data", "21/10/2026"])
    eb.main(["--data", "21/10/2026", "--substituir"])
    assert len(eb.ler_historico(eb.ARQUIVO_HISTORICO)) == 4


def test_data_invalida_da_erro_claro(projeto_ficticio):
    with pytest.raises(ValueError, match="dd/mm/aaaa"):
        eb.main(["--data", "2026-10-21"])


# --- Importação da planilha antiga ------------------------------------------

def test_importar_xlsm_cria_cadastro_e_historico(tmp_path):
    # Planilha fictícia com a mesma estrutura da Escala_Bases_Ativo2.xlsm
    wb = openpyxl.Workbook()
    escala = wb.active
    escala.title = "Escala"
    escala["A16"], escala["B16"], escala["C16"] = "Data", "BASE A", "BASE B"
    escala["A17"], escala["B17"], escala["C17"] = pd.Timestamp("2026-10-16").to_pydatetime(), "ANA", "BRUNO"
    escala["A18"], escala["B18"], escala["C18"] = pd.Timestamp("2026-10-19").to_pydatetime(), "BRUNO", "ANA"
    cadastro = wb.create_sheet("Cadastro")
    cadastro["A4"], cadastro["C4"] = "Bases", "Pessoas"
    cadastro["A5"], cadastro["C5"] = "BASE A", "ANA"
    cadastro["A6"], cadastro["C6"] = "BASE B", "BRUNO"
    antiga = tmp_path / "antiga.xlsx"
    wb.save(antiga)

    arq_cadastro = tmp_path / "config" / "bases.xlsx"
    arq_historico = tmp_path / "historico" / "historico.csv"
    eb.importar_xlsm(antiga, arq_cadastro, arq_historico)
    eb.importar_xlsm(antiga, arq_cadastro, arq_historico)  # rodar 2 vezes não duplica

    assert eb.ler_cadastro(arq_cadastro) == (["BASE A", "BASE B"], ["ANA", "BRUNO"])
    historico = eb.ler_historico(arq_historico)
    assert len(historico) == 4
    assert registro(date(2026, 10, 19), "BASE A", "BRUNO") in historico
