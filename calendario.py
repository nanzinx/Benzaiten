"""
Dias úteis e feriados.

Os feriados NACIONAIS são calculados sozinhos, inclusive a Sexta-feira Santa
(que muda todo ano porque depende da Páscoa).

Feriados estaduais e municipais, pontos facultativos (Carnaval, Corpus Christi)
e qualquer dia em que a operação não funcione ficam em config/feriados.xlsx,
preenchido por você: coluna Data (dd/mm/aaaa) e coluna Descrição.
"""

from datetime import date, datetime, timedelta
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill

from leitura import converter_data

PASTA_PROJETO = Path(__file__).resolve().parent
ARQUIVO_FERIADOS = PASTA_PROJETO / "config" / "feriados.xlsx"

# Feriados nacionais de data fixa (dia, mês, descrição)
FERIADOS_FIXOS = [
    (1, 1, "Confraternização Universal"),
    (21, 4, "Tiradentes"),
    (1, 5, "Dia do Trabalho"),
    (7, 9, "Independência do Brasil"),
    (12, 10, "Nossa Senhora Aparecida"),
    (2, 11, "Finados"),
    (15, 11, "Proclamação da República"),
    (20, 11, "Dia Nacional de Zumbi e da Consciência Negra"),
    (25, 12, "Natal"),
]


def domingo_de_pascoa(ano):
    """
    Data da Páscoa no calendário gregoriano (algoritmo de Meeus/Jones/Butcher).

    Exemplo: 2026 -> 05/04/2026
    """
    a = ano % 19
    b, c = divmod(ano, 100)
    d, e = divmod(b, 4)
    f = (b + 8) // 25
    g = (b - f + 1) // 3
    h = (19 * a + b - d - g + 15) % 30
    i, k = divmod(c, 4)
    l = (32 + 2 * e + 2 * i - h - k) % 7
    m = (a + 11 * h + 22 * l) // 451
    mes = (h + l - 7 * m + 114) // 31
    dia = (h + l - 7 * m + 114) % 31 + 1
    return date(ano, mes, dia)


def feriados_nacionais(ano):
    """{data: descrição} dos feriados nacionais do ano."""
    feriados = {date(ano, mes, dia): nome for dia, mes, nome in FERIADOS_FIXOS}
    feriados[domingo_de_pascoa(ano) - timedelta(days=2)] = "Sexta-feira Santa"
    return feriados


def ler_feriados(arquivo=None):
    """
    Feriados extras de config/feriados.xlsx: {data: descrição}.

    Se o arquivo não existir, devolve um dicionário vazio.
    """
    arquivo = Path(arquivo or ARQUIVO_FERIADOS)
    if not arquivo.exists():
        return {}
    feriados = {}
    wb = openpyxl.load_workbook(arquivo, read_only=True, data_only=True)
    for linha in wb.worksheets[0].iter_rows(min_row=2, values_only=True):
        valor, descricao = (list(linha) + [None, None])[:2]
        if isinstance(valor, datetime):
            dia = valor.date()
        elif isinstance(valor, date):
            dia = valor
        else:
            dia = converter_data(valor)
        if dia:
            feriados[dia] = str(descricao or "Feriado")
    return feriados


def criar_modelo_feriados(arquivo=None):
    """Cria config/feriados.xlsx vazio (com cabeçalho e instruções), se ainda não existir."""
    arquivo = Path(arquivo or ARQUIVO_FERIADOS)
    if arquivo.exists():
        return False
    arquivo.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Feriados"
    ws.append(["Data", "Descrição"])
    for celula in ws[1]:
        celula.font = Font(bold=True, color="FFFFFF")
        celula.fill = PatternFill("solid", fgColor="0B2A5B")
    ws["D1"] = "Os feriados NACIONAIS já são considerados automaticamente."
    ws["D2"] = "Anote aqui só os estaduais, municipais e pontos facultativos (ex.: Carnaval)."
    ws["D3"] = "Data no formato dd/mm/aaaa. Uma data por linha."
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 35
    wb.save(arquivo)
    return True


def eh_feriado(dia, extras=None):
    """Devolve a descrição do feriado (nacional ou extra) ou None."""
    extras = extras or {}
    return extras.get(dia) or feriados_nacionais(dia.year).get(dia)


def eh_dia_util(dia, extras=None):
    """Segunda a sexta, sem feriado."""
    return dia.weekday() < 5 and not eh_feriado(dia, extras)


def dia_util_anterior(hoje, extras=None):
    """
    Último dia útil antes de `hoje`.

    Segunda-feira -> sexta; depois de feriado -> o dia útil anterior a ele.
    `extras` são os feriados de config/feriados.xlsx ({data: descrição}).
    """
    dia = hoje - timedelta(days=1)
    while not eh_dia_util(dia, extras):
        dia -= timedelta(days=1)
    return dia
