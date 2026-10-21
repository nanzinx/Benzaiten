"""
Escala diária de bases do Ativo 2.

Distribui as bases (lotes/notas do INSS) entre as pessoas cadastradas
(o grupo operacional de cada supervisor) em RODÍZIO, do mesmo jeito que o botão
"Sortear bases de hoje" da planilha Escala_Bases_Ativo2.xlsm.

Regras do rodízio:
  - A cada sorteio, cada pessoa anda UMA base para frente, na ordem do cadastro
    (NOTA 0 -> NOTA 1 -> NOTA 2 -> NOTA 3 -> volta para a NOTA 0).
    Assim ninguém repete uma base antes de passar por todas.
  - O ponto de partida é o último sorteio completo do histórico (todas as bases
    preenchidas, sem ninguém repetido). Sem histórico, começa pela ordem do cadastro.
  - Pessoas ausentes no dia ficam fora; a base delas vai para quem estiver com menos
    bases (de preferência quem não teve essa base no sorteio anterior). Quando a pessoa
    volta, ela retoma a posição dela no rodízio, como se não tivesse faltado.
  - Se houver menos pessoas que bases, algumas pessoas ficam com mais de uma base.
  - Se houver mais pessoas que bases, quem fica sem base também entra no rodízio.

Como usar (no terminal, dentro da pasta do projeto):
  python escala_bases.py                          -> sorteia hoje
  python escala_bases.py --ausentes "FULANO"       -> sorteia hoje sem o FULANO
  python escala_bases.py --data 21/10/2026        -> sorteia outra data
  python escala_bases.py --importar-xlsm config/Escala_Bases_Ativo2.xlsm
                                                  -> cria config/bases.xlsx e copia o histórico da planilha antiga
"""

import argparse
import csv
import sys
from datetime import date, datetime
from pathlib import Path

import openpyxl
import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from calendario import eh_feriado, ler_feriados
from leitura import normalizar_nome

# ---------------------------------------------------------------------------
# Configurações (mude aqui se precisar)
# ---------------------------------------------------------------------------

PASTA_PROJETO = Path(__file__).resolve().parent
ARQUIVO_CADASTRO = PASTA_PROJETO / "config" / "bases.xlsx"
ARQUIVO_HISTORICO = PASTA_PROJETO / "historico" / "historico_escalas.csv"
PASTA_SAIDA = PASTA_PROJETO / "saida"

# Cores da empresa (as mesmas da planilha .xlsm)
AZUL = "0A2F6B"
AMARELO = "FFC20E"
AMARELO_CLARO = "FFF4CC"
BRANCO = "FFFFFF"

DIAS_DA_SEMANA = ["segunda-feira", "terça-feira", "quarta-feira", "quinta-feira",
                  "sexta-feira", "sábado", "domingo"]

COLUNAS_HISTORICO = ["Data", "Base", "Responsavel"]


# ---------------------------------------------------------------------------
# Leitura do cadastro e do histórico
# ---------------------------------------------------------------------------

def limpar_lista(valores):
    """Tira espaços, linhas vazias e nomes repetidos, mantendo a ordem."""
    resultado = []
    vistos = set()
    for valor in valores:
        if pd.isna(valor):
            continue
        texto = " ".join(str(valor).split())
        if texto and normalizar_nome(texto) not in vistos:
            vistos.add(normalizar_nome(texto))
            resultado.append(texto)
    return resultado


def ler_cadastro(caminho=ARQUIVO_CADASTRO):
    """
    Lê config/bases.xlsx e devolve (bases, pessoas).

    O arquivo tem duas abas:
      - "Bases":   coluna "Base"   (uma base por linha)
      - "Pessoas": coluna "Pessoa" (uma pessoa por linha)
    """
    caminho = Path(caminho)
    if not caminho.exists():
        raise FileNotFoundError(
            f"Não encontrei {caminho}.\n"
            "Crie o cadastro a partir da planilha antiga com:\n"
            "  python escala_bases.py --importar-xlsm config/Escala_Bases_Ativo2.xlsm"
        )
    bases = limpar_lista(pd.read_excel(caminho, sheet_name="Bases")["Base"])
    pessoas = limpar_lista(pd.read_excel(caminho, sheet_name="Pessoas")["Pessoa"])
    return bases, pessoas


def ler_historico(caminho=ARQUIVO_HISTORICO):
    """
    Lê o histórico acumulado e devolve uma lista de registros:
      [{"data": date, "base": "...", "responsavel": "..."}, ...]
    """
    caminho = Path(caminho)
    if not caminho.exists():
        return []
    registros = []
    with open(caminho, encoding="utf-8-sig", newline="") as arquivo:
        for linha in csv.DictReader(arquivo, delimiter=";"):
            registros.append({
                "data": datetime.strptime(linha["Data"], "%d/%m/%Y").date(),
                "base": linha["Base"],
                "responsavel": linha["Responsavel"],
            })
    return registros


def salvar_historico(registros, caminho=ARQUIVO_HISTORICO):
    """
    Grava o histórico inteiro, ordenado por data.

    Usa ; e UTF-8 com BOM para o Excel abrir com acentos certos.
    """
    caminho = Path(caminho)
    caminho.parent.mkdir(parents=True, exist_ok=True)
    # sorted é estável: dentro do mesmo dia, mantém a ordem das bases
    ordenados = sorted(registros, key=lambda r: r["data"])
    with open(caminho, "w", encoding="utf-8-sig", newline="") as arquivo:
        escritor = csv.writer(arquivo, delimiter=";")
        escritor.writerow(COLUNAS_HISTORICO)
        for r in ordenados:
            escritor.writerow([r["data"].strftime("%d/%m/%Y"), r["base"], r["responsavel"]])


# ---------------------------------------------------------------------------
# Sorteio
# ---------------------------------------------------------------------------

def ponto_de_partida(historico, bases, pessoas, data_sorteio):
    """
    Descobre a ordem do rodízio e quantos passos andar até a data do sorteio.

    Procura, do mais recente para o mais antigo, o último sorteio ANTES da data que
    esteja completo: todas as bases preenchidas e ninguém repetido (quando há pessoas
    suficientes). Dias com ausência não servem de partida, porque alguém cobriu outra base.

    Devolve (roda, passos):
      roda   = pessoas na ordem em que estavam nas bases do sorteio de partida,
               seguidas de quem não estava nele (pessoa nova no cadastro);
      passos = quantos sorteios separam a partida da data (1 = o dia seguinte).
    Sem sorteio de partida: roda = ordem do cadastro e passos = 0.
    """
    oficiais = {normalizar_nome(p): p for p in pessoas}
    datas = sorted({r["data"] for r in historico if r["data"] < data_sorteio}, reverse=True)
    for passos, d in enumerate(datas, start=1):
        do_dia = {normalizar_nome(r["base"]): normalizar_nome(r["responsavel"]) for r in historico if r["data"] == d}
        nomes = [do_dia.get(normalizar_nome(b), "") for b in bases]
        completo = all(nomes) and (len(set(nomes)) == len(nomes) or len(pessoas) < len(bases))
        if not completo:
            continue
        # Quem saiu do cadastro sai da roda; quem entrou vai para o fim dela
        roda = list(dict.fromkeys(oficiais[n] for n in nomes if n in oficiais))
        roda += [p for p in pessoas if p not in roda]
        return roda, passos
    return list(pessoas), 0


def rodizio(bases, pessoas, historico, data_sorteio, ausentes=()):
    """
    Monta a escala do dia em rodízio e devolve (escolha, avisos).

    escolha = {base: pessoa}, na ordem das bases do cadastro.
    O histórico informado não deve conter a própria data (quem chama decide se substitui).
    """
    if not bases:
        raise ValueError("Nenhuma base cadastrada.")
    fora = {normalizar_nome(a) for a in ausentes}
    presentes = [p for p in pessoas if normalizar_nome(p) not in fora]
    if not presentes:
        raise ValueError("Nenhuma pessoa disponível para o sorteio.")

    roda, passos = ponto_de_partida(historico, bases, pessoas, data_sorteio)
    m = len(roda)
    # Cada passo leva cada pessoa para a base seguinte
    plano = [roda[(j - passos) % m] for j in range(len(bases))]

    # Base de quem faltou: para quem está com menos bases; de preferência quem não teve
    # essa base no sorteio anterior; depois, quem vem logo atrás no rodízio
    anteriores = sorted({r["data"] for r in historico if r["data"] < data_sorteio})
    ontem = {normalizar_nome(r["base"]): normalizar_nome(r["responsavel"])
             for r in historico if anteriores and r["data"] == anteriores[-1]}
    escolha = {b: (p if p in presentes else None) for b, p in zip(bases, plano)}
    contagem = {p: list(escolha.values()).count(p) for p in presentes}
    avisos = []
    for base, previsto in zip(bases, plano):
        if escolha[base]:
            continue
        posicao = roda.index(previsto)
        substituto = min(presentes, key=lambda p: (
            contagem[p],
            ontem.get(normalizar_nome(base)) == normalizar_nome(p),
            (roda.index(p) - posicao) % m,
        ))
        escolha[base] = substituto
        contagem[substituto] += 1
        avisos.append(f"{previsto} está ausente: {base} ficou com {substituto}.")

    repetidos = [f"{p} ({b})" for b, p in escolha.items() if ontem.get(normalizar_nome(b)) == normalizar_nome(p)]
    if repetidos:
        avisos.append(f"por causa das ausências, repetiu a base do sorteio anterior: {', '.join(repetidos)}.")
    return escolha, avisos


def aplicar_ausencias(pessoas, ausentes):
    """
    Tira as pessoas ausentes da lista.

    Dá erro se algum nome de ausente não existir no cadastro,
    para um erro de digitação não passar despercebido.
    """
    cadastradas = {normalizar_nome(p): p for p in pessoas}
    desconhecidos = [a for a in ausentes if normalizar_nome(a) not in cadastradas]
    if desconhecidos:
        raise ValueError(
            f"Ausente(s) não encontrado(s) no cadastro: {', '.join(desconhecidos)}.\n"
            f"Pessoas cadastradas: {', '.join(pessoas)}"
        )
    fora = {normalizar_nome(a) for a in ausentes}
    return [p for p in pessoas if normalizar_nome(p) not in fora]


# ---------------------------------------------------------------------------
# Planilha de saída
# ---------------------------------------------------------------------------

def gerar_planilha(escolha, data_sorteio, ausentes, historico, caminho):
    """Cria saida/escala_AAAA-MM-DD.xlsx com a escala do dia e os últimos sorteios."""
    azul = PatternFill("solid", fgColor=AZUL)
    amarelo = PatternFill("solid", fgColor=AMARELO)
    amarelo_claro = PatternFill("solid", fgColor=AMARELO_CLARO)
    fonte_branca = Font(bold=True, color=BRANCO, size=14)
    fonte_azul = Font(bold=True, color=AZUL)
    borda = Border(*(Side(style="thin", color=AZUL),) * 4)

    wb = openpyxl.Workbook()

    # --- Aba 1: escala do dia ---
    ws = wb.active
    ws.title = "Escala"
    ws.merge_cells("A1:B1")
    ws["A1"] = "ESCALA DE BASES — ATIVO 2"
    ws["A1"].fill = azul
    ws["A1"].font = fonte_branca
    ws["A1"].alignment = Alignment(horizontal="center")
    ws.merge_cells("A2:B2")
    dia_semana = DIAS_DA_SEMANA[data_sorteio.weekday()]
    ws["A2"] = f"{data_sorteio:%d/%m/%Y} ({dia_semana})"
    ws["A2"].font = fonte_azul
    ws["A2"].alignment = Alignment(horizontal="center")

    ws.append([])
    ws.append(["Base", "Responsável"])
    for celula in ws[4]:
        celula.fill = amarelo
        celula.font = fonte_azul
        celula.border = borda
    for base, pessoa in escolha.items():
        ws.append([base, pessoa])
        for celula in ws[ws.max_row]:
            celula.border = borda

    ws.append([])
    ws.append([f"Ausentes: {', '.join(ausentes) if ausentes else 'nenhum'}"])
    ws.append(["Regra: rodízio — cada pessoa anda uma base para frente a cada sorteio."])
    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 22

    # --- Aba 2: últimos 10 sorteios, uma coluna por base (como na .xlsm) ---
    wh = wb.create_sheet("Últimos sorteios")
    bases = list(escolha)
    wh.append(["Data"] + bases)
    for celula in wh[1]:
        celula.fill = azul
        celula.font = Font(bold=True, color=BRANCO)
    datas = sorted({r["data"] for r in historico}, reverse=True)[:10]
    for d in datas:
        do_dia = {normalizar_nome(r["base"]): r["responsavel"] for r in historico if r["data"] == d}
        wh.append([d.strftime("%d/%m/%Y")] + [do_dia.get(normalizar_nome(b), "") for b in bases])
        if d == data_sorteio:
            for celula in wh[wh.max_row]:
                celula.fill = amarelo_claro
    wh.column_dimensions["A"].width = 12
    for i in range(len(bases)):
        wh.column_dimensions[openpyxl.utils.get_column_letter(i + 2)].width = 32

    caminho = Path(caminho)
    caminho.parent.mkdir(parents=True, exist_ok=True)
    wb.save(caminho)


# ---------------------------------------------------------------------------
# Importação da planilha antiga (.xlsm)
# ---------------------------------------------------------------------------

def importar_xlsm(caminho_xlsm, arquivo_cadastro=ARQUIVO_CADASTRO, arquivo_historico=ARQUIVO_HISTORICO):
    """
    Lê a planilha Escala_Bases_Ativo2.xlsm (sem alterá-la) e:
      1. cria config/bases.xlsx com as bases e pessoas da aba Cadastro
         (só se ele ainda não existir);
      2. copia para o histórico os dias da aba Escala que ainda não estão lá.
    """
    caminho_xlsm = Path(caminho_xlsm)
    if not caminho_xlsm.exists():
        raise FileNotFoundError(
            f"Não encontrei a planilha {caminho_xlsm}.\n"
            "A pasta config/ não vem junto com o projeto (ela guarda dados e fica fora do GitHub).\n"
            "Crie a pasta config dentro da pasta do projeto, copie a planilha para lá e rode de novo."
        )
    # keep_vba=False: só lemos os valores; as macros não são executadas
    wb = openpyxl.load_workbook(caminho_xlsm, read_only=True, data_only=True)

    # 1. Cadastro: bases na coluna A e pessoas na coluna C, a partir da linha 5
    cadastro = wb["Cadastro"]
    bases, pessoas = [], []
    for linha in cadastro.iter_rows(min_row=5, max_col=3, values_only=True):
        bases.append(linha[0])
        pessoas.append(linha[2])
    bases, pessoas = limpar_lista(bases), limpar_lista(pessoas)

    arquivo_cadastro = Path(arquivo_cadastro)
    if arquivo_cadastro.exists():
        print(f"- {arquivo_cadastro.name} já existe; não foi alterado.")
    else:
        arquivo_cadastro.parent.mkdir(parents=True, exist_ok=True)
        with pd.ExcelWriter(arquivo_cadastro) as escritor:
            pd.DataFrame({"Base": bases}).to_excel(escritor, sheet_name="Bases", index=False)
            pd.DataFrame({"Pessoa": pessoas}).to_excel(escritor, sheet_name="Pessoas", index=False)
        print(f"- Criado {arquivo_cadastro} com {len(bases)} bases e {len(pessoas)} pessoas.")

    # 2. Histórico: procura a linha de cabeçalho "Data | base 1 | base 2 ..."
    escala = wb["Escala"]
    linhas = list(escala.iter_rows(values_only=True))
    inicio = next(i for i, l in enumerate(linhas) if l and l[0] == "Data")
    cabecalho = linhas[inicio]

    historico = ler_historico(arquivo_historico)
    datas_existentes = {r["data"] for r in historico}
    novos = 0
    for linha in linhas[inicio + 1:]:
        if not isinstance(linha[0], datetime):
            continue
        d = linha[0].date()
        if d in datas_existentes:
            continue
        for base, pessoa in zip(cabecalho[1:], linha[1:]):
            if base and pessoa:
                historico.append({"data": d, "base": str(base).strip(), "responsavel": str(pessoa).strip()})
        datas_existentes.add(d)
        novos += 1
    salvar_historico(historico, arquivo_historico)
    print(f"- {novos} dia(s) importado(s) para {arquivo_historico}.")


# ---------------------------------------------------------------------------
# Programa principal
# ---------------------------------------------------------------------------

def confirmar(pergunta):
    """Pergunta sim/não no terminal. Sem resposta, considera 'não'."""
    try:
        return input(f"{pergunta} (s/N): ").strip().lower() in ("s", "sim")
    except EOFError:
        return False


def main(argumentos=None):
    parser = argparse.ArgumentParser(description="Sorteio diário das bases do Ativo 2.")
    parser.add_argument("--data", help="data do sorteio no formato dd/mm/aaaa (padrão: hoje)")
    parser.add_argument("--ausentes", default="", help='pessoas ausentes, separadas por vírgula. Ex.: "FULANO, BELTRANO"')
    parser.add_argument("--substituir", action="store_true", help="substitui sem perguntar se a data já foi sorteada")
    parser.add_argument("--importar-xlsm", metavar="ARQUIVO", help="importa cadastro e histórico da planilha antiga")
    args = parser.parse_args(argumentos)

    if args.importar_xlsm:
        importar_xlsm(args.importar_xlsm, ARQUIVO_CADASTRO, ARQUIVO_HISTORICO)
        return 0

    try:
        data_sorteio = datetime.strptime(args.data, "%d/%m/%Y").date() if args.data else date.today()
    except ValueError:
        raise ValueError(f"Data inválida: {args.data}. Use o formato dd/mm/aaaa, ex.: 21/10/2026")
    ausentes = [a.strip() for a in args.ausentes.split(",") if a.strip()]

    bases, pessoas = ler_cadastro(ARQUIVO_CADASTRO)
    disponiveis = aplicar_ausencias(pessoas, ausentes)
    # Mostra os ausentes com a grafia do cadastro (ex.: "fulano" -> "FULANO")
    ausentes = [p for p in pessoas if p not in disponiveis]
    historico = ler_historico(ARQUIVO_HISTORICO)

    print(f"\nSORTEIO DE BASES — ATIVO 2 — {data_sorteio:%d/%m/%Y} ({DIAS_DA_SEMANA[data_sorteio.weekday()]})")
    print(f"Bases: {len(bases)} | Pessoas disponíveis: {len(disponiveis)} | Ausentes: {', '.join(ausentes) or 'nenhum'}")

    avisos = []
    if data_sorteio.weekday() >= 5:
        avisos.append("a data escolhida cai no fim de semana.")
    feriado = eh_feriado(data_sorteio, ler_feriados())
    if feriado:
        avisos.append(f"a data escolhida é feriado ({feriado}).")
    if any(r["data"] == data_sorteio for r in historico):
        if not args.substituir and not confirmar(f"Já existe sorteio para {data_sorteio:%d/%m/%Y}. Substituir?"):
            print("Nada foi alterado.")
            return 1
        historico = [r for r in historico if r["data"] != data_sorteio]
    if len(disponiveis) < len(bases):
        avisos.append(f"{len(disponiveis)} pessoa(s) para {len(bases)} base(s): alguém fica com mais de uma base.")

    escolha, avisos_rodizio = rodizio(bases, pessoas, historico, data_sorteio, ausentes)
    avisos += avisos_rodizio
    sem_base = [p for p in disponiveis if p not in escolha.values()]
    if sem_base:
        avisos.append(f"mais pessoas que bases; sem base hoje: {', '.join(sem_base)}. "
                      "Cadastre novas bases em config/bases.xlsx se precisar.")

    # Grava histórico e planilha (arquivos novos; o cadastro não é alterado)
    for base, pessoa in escolha.items():
        historico.append({"data": data_sorteio, "base": base, "responsavel": pessoa})
    salvar_historico(historico, ARQUIVO_HISTORICO)
    saida = PASTA_SAIDA / f"escala_{data_sorteio:%Y-%m-%d}.xlsx"
    gerar_planilha(escolha, data_sorteio, ausentes, historico, saida)

    print()
    largura = max(len(b) for b in escolha)
    for base, pessoa in escolha.items():
        print(f"  {base:<{largura}}  =  {pessoa}")
    print()
    for aviso in avisos:
        print(f"AVISO: {aviso}")
    print(f"Planilha: {saida}")
    print(f"Histórico: {ARQUIVO_HISTORICO}")
    return 0


if __name__ == "__main__":
    try:
        sys.exit(main())
    except PermissionError as erro:
        print(f"\nERRO: não consegui gravar {erro.filename}.\n"
              "Se esse arquivo estiver aberto no Excel, feche-o e rode de novo.")
        sys.exit(1)
    except (FileNotFoundError, ValueError) as erro:
        print(f"\nERRO: {erro}")
        sys.exit(1)
