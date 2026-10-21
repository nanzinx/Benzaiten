"""
Ranking de contratos pagos (vendedores e supervisores).

Substitui a macro ATUALIZAR da planilha Ranking_relatorio_com_comentarios_OFICIAL.xlsm.
A macro lia a aba DADOS_BRUTOS pela POSIÇÃO das colunas (J, S, T, U, Y) e quebrou
quando a esteira nova ganhou uma coluna. Aqui as colunas são lidas pelo NOME.

Entradas (em entrada/, reconhecidas pelas colunas):
  - esteira de "data de andamento" e esteira de "data de pagamento" do Vanguard
    (as duas são juntadas e cada proposta conta uma vez só)
Configuração:
  - config/operadores.xlsx  vendedores ativos e apelidos (o mesmo do relatório da URA)
  - config/ranking.xlsx     metas, vendedores CLT, supervisores CLT e nomes das equipes

Colunas do ranking (mesmas regras da macro):
  PORTABILIDADE ....... soma do Valor Contrato das propostas PAGAS no mês
                        (sem margem livre e cartão)
  PROJETANDO .......... portabilidade / esperado para a data
                        (esperado = meta / dias úteis do mês x dias úteis até a data)
  POTENCIAL DE GANHOS . (bônus da faixa + 1% da margem/cartão + extra) x 1,2
  MARGEM/CARTÃO ....... quantidade de propostas pagas de margem livre ou cartão
  ANDAMENTO ........... soma das propostas em andamento (sem margem livre e cartão)
  PORT. PROJEÇÃO ...... portabilidade + metade do andamento
  ELEGÍVEL AO BÔNUS ... SIM se a portabilidade atingiu a meta do bônus

Como usar:
  python ranking_vendas.py                    -> ranking com dados até o último dia útil
  python ranking_vendas.py --data 20/10/2026  -> ranking com dados até outra data
  python ranking_vendas.py --importar-xlsm config/Ranking_relatorio_com_comentarios_OFICIAL.xlsm
                                              -> cria config/ranking.xlsx a partir da planilha antiga
"""

import argparse
import sys
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

import relatorio_ura as ru
from calendario import dia_util_anterior, eh_dia_util, ler_feriados
from esteira import eh_cartao, nome_do_agente, remover_propostas_repetidas
from leitura import converter_data, converter_valor_br, normalizar_nome

# ---------------------------------------------------------------------------
# Configurações (regras de bônus copiadas da macro original)
# ---------------------------------------------------------------------------

PASTA_PROJETO = Path(__file__).resolve().parent
PASTA_ENTRADA = PASTA_PROJETO / "entrada"
PASTA_SAIDA = PASTA_PROJETO / "saida"
ARQUIVO_CONFIG = PASTA_PROJETO / "config" / "ranking.xlsx"

COLUNA_VALOR = "VALOR CONTRATO"
ETAPA_PAGO = "PAGO"
ETAPAS_ANDAMENTO = {"ANDAMENTO"}   # acrescente "PENDENTE" aqui se ela também contar

# Faixas do bônus: (atingimento até, taxa sobre a meta x atingimento).
# A última faixa vale para qualquer atingimento acima da anterior.
FAIXAS_VENDEDOR = {
    "CLT": [(1.0, 0.0015), (550 / 450, 0.0025), (float("inf"), 0.003)],
    "ESTAGIO": [(1.0, 0.001), (500 / 350, 0.0025), (float("inf"), 0.003)],
}
FAIXAS_SUPERVISOR = {
    "CLT": [(1.0, 0.001), (3600 / 2800, 0.0015), (float("inf"), 0.002)],
    "ESTAGIO": [(1.0, 0.0015), (2800 / 2400, 0.0018), (float("inf"), 0.0024)],
}
BONUS_EXTRA_VENDEDOR = 1000   # na última faixa (só vendedores)
TAXA_MARGEM_CARTAO = 0.01     # 1% do valor pago em margem livre/cartão
ACRESCIMO_BONUS = 1.2         # "bônus de 20%" sobre o total
PESO_ANDAMENTO = 0.5          # projeção: metade do andamento (a outra metade reprova)
DESTAQUE = 1.33               # a partir de 133% do esperado, o nome ganha o emoji
COR_ACIMA_DA_META = "1F4EBF"  # projetando >= 100%: azul
COR_PERTO_DA_META = "00A651"  # projetando >= 70%: verde

AZUL = ru.AZUL
AMARELO = ru.AMARELO
BRANCO = ru.BRANCO
PRETO = "000000"
VERDE_SIM = "00B050"
VERMELHO_NAO = "C00000"

TITULOS = ["RANKING", "NOME", "PORTABILIDADE", "PROJETANDO", "POTENCIAL DE GANHOS",
           "MARGEM/CARTÃO", "ANDAMENTO", "PORT. PROJEÇÃO", "ELEGÍVEL AO BÔNUS", "FEEDBACK"]


# ---------------------------------------------------------------------------
# Configuração (config/ranking.xlsx)
# ---------------------------------------------------------------------------

# Parâmetro na planilha -> chave usada no código
PARAMETROS = {
    "META DO VENDEDOR": "meta_vendedor",
    "META DO VENDEDOR PARA O BONUS": "bonus_vendedor",
    "META DO SUPERVISOR": "meta_supervisor",
    "META DO SUPERVISOR PARA O BONUS": "bonus_supervisor",
    "META GERAL": "meta_geral",
    "META DE CARTAO": "meta_cartao",
    "EMOJI DE DESTAQUE": "emoji",
}


@dataclass
class Config:
    """Metas (por categoria CLT/ESTAGIO), listas e nomes das equipes."""
    metas: dict                                     # {chave: {"CLT": valor, "ESTAGIO": valor}}
    vendedores_clt: set = field(default_factory=set)      # nomes normalizados
    supervisores_clt: set = field(default_factory=set)    # equipes normalizadas
    nomes_equipes: dict = field(default_factory=dict)     # {EQUIPE NORMALIZADA: nome no relatório}

    def meta(self, chave, categoria="CLT"):
        valores = self.metas.get(chave) or {}
        valor = valores.get(categoria)
        return valores.get("CLT") if valor in (None, "") else valor


def salvar_config(arquivo, metas, vendedores_clt, supervisores_clt, nomes_equipes):
    """Grava config/ranking.xlsx com 4 abas, todas com cabeçalho e explicação."""
    arquivo = Path(arquivo)
    arquivo.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()

    def cabecalho(ws, titulos, larguras, nota):
        ws.append(titulos)
        for celula in ws[1]:
            celula.font = Font(bold=True, color=BRANCO)
            celula.fill = PatternFill("solid", fgColor=AZUL)
        for i, largura in enumerate(larguras):
            ws.column_dimensions[openpyxl.utils.get_column_letter(i + 1)].width = largura
        ws.cell(1, len(titulos) + 2, nota).font = Font(italic=True, color=AZUL)

    ws = wb.active
    ws.title = "Parametros"
    cabecalho(ws, ["Parâmetro", "CLT", "Estágio"], [36, 16, 16],
              "Coluna Estágio: vendedores fora da aba 'Vendedores CLT' e supervisores fora de 'Supervisores CLT'.")
    nomes = {chave: rotulo for rotulo, chave in PARAMETROS.items()}
    rotulos_bonitos = {
        "meta_vendedor": "Meta do vendedor", "bonus_vendedor": "Meta do vendedor para o bônus",
        "meta_supervisor": "Meta do supervisor", "bonus_supervisor": "Meta do supervisor para o bônus",
        "meta_geral": "Meta geral", "meta_cartao": "Meta de cartão", "emoji": "Emoji de destaque",
    }
    for chave in nomes:
        valores = metas.get(chave, {})
        ws.append([rotulos_bonitos[chave], valores.get("CLT"), valores.get("ESTAGIO")])

    for titulo, coluna, itens, nota in [
        ("Vendedores CLT", "Nome", sorted(vendedores_clt),
         "Vendedores CLT (todos os outros usam as metas de Estágio). Mantenha atualizado!"),
        ("Supervisores CLT", "Equipe", sorted(supervisores_clt),
         "Equipes cujo supervisor usa as metas da coluna CLT."),
    ]:
        aba = wb.create_sheet(titulo)
        cabecalho(aba, [coluna], [45], nota)
        for item in itens:
            aba.append([item])

    aba = wb.create_sheet("Nomes das equipes")
    cabecalho(aba, ["Equipe na esteira", "Nome no relatório"], [40, 40],
              "Opcional: troca o nome da equipe no ranking dos supervisores.")
    for de, para in nomes_equipes.items():
        aba.append([de, para])
    wb.save(arquivo)


def ler_config(arquivo=None):
    """Lê config/ranking.xlsx. Dá erro explicando como criar se não existir."""
    arquivo = Path(arquivo or ARQUIVO_CONFIG)
    if not arquivo.exists():
        raise FileNotFoundError(
            f"Não encontrei {arquivo}.\n"
            "Crie a partir da planilha antiga com:\n"
            "  python ranking_vendas.py --importar-xlsm config/Ranking_relatorio_com_comentarios_OFICIAL.xlsm"
        )
    wb = openpyxl.load_workbook(arquivo, read_only=True, data_only=True)
    metas = {}
    for rotulo, clt, estagio in wb["Parametros"].iter_rows(min_row=2, max_col=3, values_only=True):
        chave = PARAMETROS.get(normalizar_nome(rotulo))
        if chave:
            metas[chave] = {"CLT": clt, "ESTAGIO": estagio}

    def coluna_a(nome_aba):
        if nome_aba not in wb.sheetnames:
            return set()
        return {normalizar_nome(v) for (v, *_) in wb[nome_aba].iter_rows(min_row=2, values_only=True) if v}

    nomes_equipes = {}
    if "Nomes das equipes" in wb.sheetnames:
        for linha in wb["Nomes das equipes"].iter_rows(min_row=2, max_col=2, values_only=True):
            de, para = (list(linha) + [None, None])[:2]
            if de and para:
                nomes_equipes[normalizar_nome(de)] = str(para).strip()
    faltando = [r for r, c in PARAMETROS.items() if c != "emoji" and not (metas.get(c) or {}).get("CLT")]
    if faltando:
        raise ValueError(f"Preencha em {arquivo} (aba Parametros): {', '.join(faltando)}")
    return Config(metas, coluna_a("Vendedores CLT"), coluna_a("Supervisores CLT"), nomes_equipes)


def importar_xlsm(caminho_xlsm, arquivo_config=None):
    """
    Cria config/ranking.xlsx a partir das abas GUIA e AUXILIAR da planilha antiga
    (sem alterá-la e sem executar as macros).
    """
    caminho_xlsm = Path(caminho_xlsm)
    if not caminho_xlsm.exists():
        raise FileNotFoundError(f"Não encontrei a planilha {caminho_xlsm}.")
    wb = openpyxl.load_workbook(caminho_xlsm, read_only=True, data_only=True)
    guia, aux = wb["GUIA"], wb["AUXILIAR"]
    g = lambda celula: guia[celula].value
    metas = {
        "meta_vendedor": {"CLT": g("B8"), "ESTAGIO": g("E8")},
        "bonus_vendedor": {"CLT": g("B10"), "ESTAGIO": g("E10")},
        "meta_supervisor": {"CLT": g("B14"), "ESTAGIO": g("E14")},
        "bonus_supervisor": {"CLT": g("B11"), "ESTAGIO": g("E11")},
        "meta_geral": {"CLT": g("B17"), "ESTAGIO": None},
        "meta_cartao": {"CLT": g("B7"), "ESTAGIO": g("E7")},
        "emoji": {"CLT": g("B19"), "ESTAGIO": None},
    }
    linhas = list(aux.iter_rows(min_row=1, max_col=3, values_only=True))
    # Vendedores CLT: coluna A, da linha 14 até a primeira vazia
    vendedores = []
    for a, _b, _c in linhas[13:]:
        if not a:
            break
        vendedores.append(" ".join(str(a).split()))
    # Supervisores CLT: coluna C, abaixo do título "Tabela de Supervisor CLT"
    supervisores = [" ".join(str(c).split()) for _a, _b, c in linhas[1:] if c]
    # Troca de nomes das equipes: A -> B nas primeiras linhas (até a instrução "ESSE TRECHO...")
    nomes = {}
    for a, b, _c in linhas[:6]:
        if a and b:
            nomes[" ".join(str(a).split())] = " ".join(str(b).split())

    arquivo_config = Path(arquivo_config or ARQUIVO_CONFIG)
    if arquivo_config.exists():
        print(f"- {arquivo_config} já existe; não foi alterado.")
        return False
    salvar_config(arquivo_config, metas, vendedores, supervisores, nomes)
    print(f"- Criado {arquivo_config}: {len(vendedores)} vendedor(es) CLT, "
          f"{len(supervisores)} supervisor(es) CLT, {len(nomes)} troca(s) de nome de equipe.")
    return True


# ---------------------------------------------------------------------------
# Cálculos
# ---------------------------------------------------------------------------

def dias_uteis(inicio, fim, feriados):
    """Quantos dias úteis existem de `inicio` a `fim` (os dois inclusos)."""
    total, dia = 0, inicio
    while dia <= fim:
        total += eh_dia_util(dia, feriados)
        dia += timedelta(days=1)
    return total


def fim_do_mes(dia):
    proximo = (dia.replace(day=28) + timedelta(days=4)).replace(day=1)
    return proximo - timedelta(days=1)


@dataclass
class Linha:
    """Uma linha do ranking (vendedor ou supervisor)."""
    nome: str
    categoria: str = "ESTAGIO"     # CLT ou ESTAGIO
    portabilidade: float = 0.0
    qtd_cartao: int = 0
    valor_cartao: float = 0.0
    andamento: float = 0.0
    projetando: float = 0.0
    potencial: float = 0.0
    elegivel: bool = False
    posicao: int = 0
    feedback: str = ""

    @property
    def projecao(self):
        return self.portabilidade + self.andamento * PESO_ANDAMENTO


@dataclass
class Ranking:
    data: date
    esperado_fator: float                 # dias úteis até a data / dias úteis do mês
    dias_mes: int
    dias_ate: int
    vendedores: list
    supervisores: list
    total: dict
    avisos: list = field(default_factory=list)
    arquivos: list = field(default_factory=list)
    usados: list = field(default_factory=list)
    ignorados: dict = field(default_factory=dict)   # {agente: valor pago ignorado}
    contagem: dict = field(default_factory=dict)


def bonus_da_faixa(atingimento, meta, faixas):
    """Taxa da faixa aplicada sobre meta x atingimento. Devolve (bônus, está na última faixa)."""
    for i, (limite, taxa) in enumerate(faixas):
        if atingimento < limite or i == len(faixas) - 1:
            return meta * atingimento * taxa, i == len(faixas) - 1
    return 0.0, False


def classificar(linhas):
    """Ordena pelo projetando (empate: portabilidade) e numera; empates exatos dividem a posição."""
    linhas.sort(key=lambda l: (-l.projetando, -l.portabilidade, l.nome))
    for l in linhas:
        l.posicao = 1 + sum(1 for outro in linhas if outro.projetando > l.projetando)
    return linhas


def moeda_curta(valor):
    """1234567.8 -> 'R$ 1.234.568'"""
    return f"R$ {valor:,.0f}".replace(",", ".")


PALAVRAS_DE_EQUIPE = {"SEM", "EQUIPE", "TIME", "ADMINISTRATIVO"}


def como_chamar(nome):
    """
    Como tratar a pessoa no feedback: o primeiro nome ('JOANA DE SOUZA' -> 'Joana').
    Nomes de equipe ('SEM SUPERVISOR', 'FULANO - URA') ficam completos.
    """
    partes = nome.split()
    if not partes or "-" in nome or normalizar_nome(partes[0]) in PALAVRAS_DE_EQUIPE:
        return " ".join(p.capitalize() for p in partes) or nome
    return partes[0].capitalize()


def gerar_feedback(linha, esperado):
    """Comentário local, objetivo e respeitoso, montado a partir dos números (sem internet)."""
    quem = como_chamar(linha.nome)
    pct = f"{linha.projetando:.0%}"
    falta = moeda_curta(max(esperado - linha.portabilidade, 0))
    if linha.projetando >= DESTAQUE:
        return f"{quem}, resultado excepcional: {pct} do esperado. Parabéns!"
    if linha.projetando >= 1:
        return f"{quem}, esperado do período atingido ({pct}). Continue nesse ritmo!"
    if linha.projetando >= 0.7:
        return f"{quem}, faltam {falta} para o esperado. Dá para chegar!"
    if linha.portabilidade > 0 and linha.andamento > 0:
        return (f"{quem}, {pct} do esperado. As propostas em andamento "
                f"({moeda_curta(linha.andamento)}) podem virar o jogo.")
    if linha.portabilidade > 0:
        return f"{quem}, {pct} do esperado; faltam {falta}. Vamos reforçar a esteira!"
    if linha.andamento > 0:
        return f"{quem}, ainda sem pagos no mês, com {moeda_curta(linha.andamento)} em andamento. Acompanhe de perto!"
    return f"{quem}, ainda sem propostas no mês. Conte com a supervisão para começar."


def calcular(dia, entradas, ativos, inativos, apelidos, config, feriados):
    """Monta o ranking de vendedores e supervisores com os dados até `dia`."""
    inicio, fim = dia.replace(day=1), fim_do_mes(dia)
    dias_mes = dias_uteis(inicio, fim, feriados)
    dias_ate = dias_uteis(inicio, dia, feriados)
    fator = dias_ate / dias_mes if dias_mes else 0

    avisos, arquivos, usados, contagem = [], [], [], {}
    todas = []
    for arquivo, linhas in entradas["esteira"]:
        arquivos.append((arquivo.name, f"{len(linhas)} proposta(s)"))
        usados.append(arquivo)
        todas += linhas
    if not todas:
        raise FileNotFoundError("Não há esteira do Vanguard em entrada/ (de andamento e de pagamento).")
    propostas, repetidas = remover_propostas_repetidas(todas)
    contagem["Propostas lidas"] = len(todas)
    contagem["Propostas repetidas entre as esteiras (contadas uma vez)"] = repetidas

    # Separa o que entra no ranking
    pagos, andamento = [], []
    pagos_fora, pagos_sem_data, outras_etapas = 0, 0, {}
    for p in propostas:
        etapa = normalizar_nome(p.get("ETAPA"))
        if etapa == ETAPA_PAGO:
            pago_em = converter_data(p.get("DATA PAGAMENTO"))
            if pago_em is None:
                pagos_sem_data += 1
                pagos.append(p)
            elif inicio <= pago_em <= fim:
                pagos.append(p)
            else:
                pagos_fora += 1
        elif etapa in ETAPAS_ANDAMENTO:
            andamento.append(p)
        else:
            outras_etapas[p.get("ETAPA") or "(vazia)"] = outras_etapas.get(p.get("ETAPA") or "(vazia)", 0) + 1
    contagem["Pagas no mês"] = len(pagos)
    contagem["Em andamento"] = len(andamento)
    contagem["Pagas fora do mês (ignoradas)"] = pagos_fora
    for etapa, qtd in sorted(outras_etapas.items()):
        contagem[f"Etapa '{etapa}' (não entra no ranking)"] = qtd
    if pagos_sem_data:
        avisos.append(f"{pagos_sem_data} proposta(s) pagas sem Data Pagamento entraram no mês por não terem data.")
    if outras_etapas:
        avisos.append("Etapas que não entram no ranking: " + ", ".join(f"{e} ({q})" for e, q in outras_etapas.items())
                      + ". Se alguma deve contar como andamento, inclua em ETAPAS_ANDAMENTO.")

    # --- Vendedores (só os ativos no cadastro) ---
    vendedores, ignorados = {}, {}

    def vendedor(texto):
        chave = normalizar_nome(nome_do_agente(texto))
        chave = apelidos.get(chave, chave)
        return chave if chave in ativos else None

    for p in pagos + andamento:
        if eh_cartao(p):
            continue  # como na macro: só portabilidade cria a linha do vendedor
        chave = vendedor(p.get("AGENTE"))
        if chave is None:
            nome = nome_do_agente(p.get("AGENTE")) or "(sem agente)"
            if normalizar_nome(p.get("ETAPA")) == ETAPA_PAGO:
                ignorados[nome] = ignorados.get(nome, 0) + converter_valor_br(p.get(COLUNA_VALOR))
            else:
                ignorados.setdefault(nome, 0)
            continue
        if chave not in vendedores:
            nome = ativos[chave][0]
            categoria = "CLT" if chave in config.vendedores_clt else "ESTAGIO"
            vendedores[chave] = Linha(nome, categoria)
    for p in pagos + andamento:
        chave = vendedor(p.get("AGENTE"))
        if chave not in vendedores:
            continue
        valor = converter_valor_br(p.get(COLUNA_VALOR))
        linha = vendedores[chave]
        pago = normalizar_nome(p.get("ETAPA")) == ETAPA_PAGO
        if eh_cartao(p):
            if pago:
                linha.qtd_cartao += 1
                linha.valor_cartao += valor
        elif pago:
            linha.portabilidade += valor
        else:
            linha.andamento += valor

    for linha in vendedores.values():
        meta = config.meta("meta_vendedor", linha.categoria)
        esperado = meta * fator
        linha.projetando = linha.portabilidade / esperado if esperado else 0
        bonus, ultima = bonus_da_faixa(linha.projetando, meta, FAIXAS_VENDEDOR[linha.categoria])
        extra = BONUS_EXTRA_VENDEDOR if ultima else 0
        linha.potencial = (bonus + linha.valor_cartao * TAXA_MARGEM_CARTAO + extra) * ACRESCIMO_BONUS
        linha.elegivel = linha.portabilidade >= config.meta("bonus_vendedor", linha.categoria)
        linha.feedback = gerar_feedback(linha, esperado)

    # --- Supervisores (por equipe da esteira; todas as propostas, como na macro) ---
    supervisores = {}
    for p in pagos + andamento:
        equipe = " ".join(str(p.get("EQUIPE") or "").split())
        if not equipe:
            continue
        chave = normalizar_nome(nome_do_agente(equipe))
        if chave not in supervisores:
            if eh_cartao(p):
                continue
            nome = config.nomes_equipes.get(normalizar_nome(equipe), equipe)
            categoria = "CLT" if normalizar_nome(equipe) in config.supervisores_clt else "ESTAGIO"
            supervisores[chave] = Linha(nome, categoria)
    for p in pagos + andamento:
        chave = normalizar_nome(nome_do_agente(" ".join(str(p.get("EQUIPE") or "").split())))
        if chave not in supervisores:
            continue
        linha = supervisores[chave]
        valor = converter_valor_br(p.get(COLUNA_VALOR))
        pago = normalizar_nome(p.get("ETAPA")) == ETAPA_PAGO
        if eh_cartao(p):
            if pago:
                linha.qtd_cartao += 1
                linha.valor_cartao += valor
        elif pago:
            linha.portabilidade += valor
        else:
            linha.andamento += valor
    for linha in supervisores.values():
        meta = config.meta("meta_supervisor", linha.categoria)
        esperado = meta * fator
        linha.projetando = linha.portabilidade / esperado if esperado else 0
        bonus, _ = bonus_da_faixa(linha.projetando, meta, FAIXAS_SUPERVISOR[linha.categoria])
        linha.potencial = (bonus + linha.valor_cartao * TAXA_MARGEM_CARTAO) * ACRESCIMO_BONUS
        linha.elegivel = linha.portabilidade >= config.meta("bonus_supervisor", linha.categoria)
        linha.feedback = gerar_feedback(linha, esperado)

    vend = classificar(list(vendedores.values()))
    sup = classificar(list(supervisores.values()))

    # Total geral: soma dos supervisores (potencial e elegíveis somam as duas tabelas)
    portabilidade = sum(l.portabilidade for l in sup)
    esperado_geral = config.meta("meta_geral") * fator
    total = {
        "portabilidade": portabilidade,
        "projetando": portabilidade / esperado_geral if esperado_geral else 0,
        "potencial": sum(l.potencial for l in vend + sup),
        "qtd_cartao": sum(l.qtd_cartao for l in sup),
        "andamento": sum(l.andamento for l in sup),
        "projecao": sum(l.projecao for l in sup),
        "elegiveis": sum(1 for l in vend + sup if l.elegivel),
    }
    ranking = Ranking(dia, fator, dias_mes, dias_ate, vend, sup, total, avisos, arquivos, usados,
                      ignorados, contagem)
    if ignorados:
        avisos.append(f"{len(ignorados)} agente(s) fora do cadastro de ativos ficaram fora do ranking "
                      "(veja a aba Conferência).")
    return ranking


# ---------------------------------------------------------------------------
# Planilha de saída
# ---------------------------------------------------------------------------

def preencher(celula, fundo=None, fonte=None, negrito=False, tamanho=11, alinhar=None):
    if fundo:
        celula.fill = PatternFill("solid", fgColor=fundo)
    celula.font = Font(bold=negrito, color=fonte, size=tamanho)
    if alinhar:
        celula.alignment = Alignment(horizontal=alinhar, vertical="center", wrap_text=True)


def gerar_planilha(ranking, config, caminho, hoje):
    """Cria o .xlsx com a aba Ranking (vendedores, supervisores, total, legenda) e a Conferência."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Ranking"
    ultima = 2 + len(TITULOS)          # B..K = tabela; A e L = faixas azuis
    emoji = config.meta("emoji") or ""
    projecao_em = hoje + timedelta(days=7)
    periodo = f"DE {ranking.data.replace(day=1):%d/%m/%Y} A {ranking.data:%d/%m/%Y}"
    linha_excel = 2

    def faixa_azul(n):
        preencher(ws.cell(n, 1), AZUL)
        preencher(ws.cell(n, ultima), AZUL)

    def titulo(texto):
        nonlocal linha_excel
        ws.merge_cells(start_row=linha_excel, start_column=2, end_row=linha_excel, end_column=ultima - 1)
        celula = ws.cell(linha_excel, 2, texto)
        preencher(celula, AMARELO, PRETO, True, 16, "center")
        ws.row_dimensions[linha_excel].height = 30
        faixa_azul(linha_excel)
        linha_excel += 1
        for i, t in enumerate(TITULOS):
            rotulo = f"{t} - {projecao_em:%d/%m/%Y}" if t == "PORT. PROJEÇÃO" else t
            preencher(ws.cell(linha_excel, 2 + i, rotulo), AZUL, BRANCO, True, 11, "center")
        ws.row_dimensions[linha_excel].height = 32
        faixa_azul(linha_excel)
        linha_excel += 1

    def tabela(linhas):
        nonlocal linha_excel
        for l in linhas:
            nome = f"{l.nome}   {emoji}" if l.projetando >= DESTAQUE and emoji else l.nome
            cor = COR_ACIMA_DA_META if l.projetando >= 1 else COR_PERTO_DA_META if l.projetando >= 0.7 else PRETO
            valores = [f"{l.posicao}º", nome, l.portabilidade, l.projetando, l.potencial, l.qtd_cartao,
                       l.andamento, l.projecao, "SIM" if l.elegivel else "NÃO", l.feedback]
            formatos = [None, None, '"R$ "#,##0.00', "0.0%", '"R$ "#,##0.00', "0", '"R$ "#,##0.00',
                        '"R$ "#,##0.00', None, None]
            for i, (valor, formato) in enumerate(zip(valores, formatos)):
                celula = ws.cell(linha_excel, 2 + i, valor)
                preencher(celula, None, cor, True, 12, "left" if i in (1, 9) else "center")
                if formato:
                    celula.number_format = formato
            preencher(ws.cell(linha_excel, 2), AMARELO, PRETO, True, 12, "center")
            elegivel = ws.cell(linha_excel, 10)
            preencher(elegivel, VERDE_SIM if l.elegivel else VERMELHO_NAO, BRANCO, True, 12, "center")
            faixa_azul(linha_excel)
            linha_excel += 1

    def separador():
        nonlocal linha_excel
        for c in range(1, ultima + 1):
            preencher(ws.cell(linha_excel, c), AZUL)
        linha_excel += 1

    titulo(f"RANKING GERAL DE CONTRATOS PAGOS - VENDEDOR - {periodo}")
    tabela(ranking.vendedores)
    separador()
    titulo(f"RANKING GERAL DE CONTRATOS PAGOS - SUPERVISOR - {periodo}")
    tabela(ranking.supervisores)

    # Total geral
    t = ranking.total
    valores = ["-", "Total Geral", t["portabilidade"], t["projetando"], t["potencial"], t["qtd_cartao"],
               t["andamento"], t["projecao"], t["elegiveis"], ""]
    formatos = [None, None, '"R$ "#,##0.00', "0.0%", '"R$ "#,##0.00', "0", '"R$ "#,##0.00', '"R$ "#,##0.00', "0", None]
    for i, (valor, formato) in enumerate(zip(valores, formatos)):
        celula = ws.cell(linha_excel, 2 + i, valor)
        preencher(celula, PRETO, BRANCO, True, 13, "center")
        if formato:
            celula.number_format = formato
    faixa_azul(linha_excel)
    linha_excel += 1
    separador()

    # Legenda
    legenda = [
        ("LEGENDA", ""),
        ("POTENCIAL DE GANHOS", "Quanto o vendedor tem a ganhar mantendo o desempenho atual."),
        ("ANDAMENTO", "Valor total de contratos em andamento, ou seja, que ainda serão pagos."),
        ("PROJEÇÃO", "Portabilidade paga + metade da esteira em andamento (a outra metade costuma reprovar)."),
        ("PROJETANDO", f"Pago ÷ esperado para a data ({ranking.dias_ate} de {ranking.dias_mes} dias úteis do mês). "
                       "Azul: 100% ou mais; verde: 70% ou mais."),
        ("META DE CARTÃO", f"{config.meta('meta_cartao')} por vendedor"),
    ]
    for rotulo, texto in legenda:
        ws.merge_cells(start_row=linha_excel, start_column=4, end_row=linha_excel, end_column=5)
        ws.merge_cells(start_row=linha_excel, start_column=6, end_row=linha_excel, end_column=ultima - 1)
        preencher(ws.cell(linha_excel, 4, rotulo), PRETO, BRANCO, True, 11, "center")
        preencher(ws.cell(linha_excel, 6, texto), PRETO, BRANCO, False, 11, "left")
        faixa_azul(linha_excel)
        linha_excel += 1
    separador()

    larguras = [2, 10, 44, 18, 13, 18, 14, 18, 22, 14, 70, 2]
    for i, largura in enumerate(larguras):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i + 1)].width = largura
    ws.freeze_panes = "D4"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    # --- Conferência ---
    wc = wb.create_sheet("Conferência")
    wc.column_dimensions["A"].width = 60
    wc.column_dimensions["B"].width = 22
    wc.column_dimensions["C"].width = 22

    def secao(texto):
        wc.append([])
        wc.append([texto])
        preencher(wc.cell(wc.max_row, 1), AZUL, BRANCO, True)

    wc.append([f"Conferência do ranking — dados até {ranking.data:%d/%m/%Y}"])
    wc.cell(1, 1).font = Font(bold=True, size=13, color=AZUL)
    secao("Dias úteis e esperado para a data")
    wc.append(["Dias úteis do mês", ranking.dias_mes])
    wc.append(["Dias úteis até a data", ranking.dias_ate])
    wc.append(["", "CLT", "Estágio"])
    for rotulo, chave in [("Esperado - vendedor", "meta_vendedor"), ("Esperado - supervisor", "meta_supervisor")]:
        wc.append([rotulo, config.meta(chave, "CLT") * ranking.esperado_fator,
                   config.meta(chave, "ESTAGIO") * ranking.esperado_fator])
        for c in (2, 3):
            wc.cell(wc.max_row, c).number_format = '"R$ "#,##0.00'
    wc.append(["Esperado - total geral", config.meta("meta_geral") * ranking.esperado_fator])
    wc.cell(wc.max_row, 2).number_format = '"R$ "#,##0.00'
    secao("Arquivos lidos")
    for nome, obs in ranking.arquivos:
        wc.append([nome, obs])
    secao("Propostas")
    for rotulo, qtd in ranking.contagem.items():
        wc.append([rotulo, qtd])
    secao("Agentes fora do cadastro de ativos (ignorados) — valor pago")
    for nome, valor in sorted(ranking.ignorados.items()):
        wc.append([nome, round(valor, 2)])
        wc.cell(wc.max_row, 2).number_format = '"R$ "#,##0.00'
    secao("Avisos")
    for aviso in ranking.avisos:
        wc.append([aviso])

    caminho = Path(caminho)
    caminho.parent.mkdir(parents=True, exist_ok=True)
    wb.save(caminho)


# ---------------------------------------------------------------------------
# Programa principal
# ---------------------------------------------------------------------------

def main(argumentos=None, usados=None):
    """
    `usados`: se for uma lista, recebe os arquivos de entrada/ usados no ranking
    (o benzaiten.py usa isso para arquivar tudo no final).
    """
    parser = argparse.ArgumentParser(description="Ranking de contratos pagos (vendedores e supervisores).")
    parser.add_argument("--data", help="dados até esta data dd/mm/aaaa (padrão: último dia útil)")
    parser.add_argument("--importar-xlsm", metavar="ARQUIVO", help="cria config/ranking.xlsx a partir da planilha antiga")
    parser.add_argument("--entrada", default=str(PASTA_ENTRADA), help=argparse.SUPPRESS)
    parser.add_argument("--saida", default=str(PASTA_SAIDA), help=argparse.SUPPRESS)
    parser.add_argument("--cadastro", default=str(ru.ARQUIVO_CADASTRO), help=argparse.SUPPRESS)
    parser.add_argument("--config", default=str(ARQUIVO_CONFIG), help=argparse.SUPPRESS)
    parser.add_argument("--hoje", help=argparse.SUPPRESS)  # usado nos testes
    args = parser.parse_args(argumentos)

    if args.importar_xlsm:
        importar_xlsm(args.importar_xlsm, args.config)
        return 0

    hoje = datetime.strptime(args.hoje, "%d/%m/%Y").date() if args.hoje else date.today()
    feriados = ler_feriados()
    if args.data:
        try:
            dia = datetime.strptime(args.data, "%d/%m/%Y").date()
        except ValueError:
            raise ValueError(f"Data inválida: {args.data}. Use o formato dd/mm/aaaa, ex.: 20/10/2026")
    else:
        dia = dia_util_anterior(hoje, feriados)

    config = ler_config(args.config)
    entradas, _ = ru.ler_entradas(args.entrada)
    if entradas["operadores"]:
        arquivo, linhas = max(entradas["operadores"], key=lambda par: par[0].stat().st_mtime)
        ru.atualizar_cadastro(linhas, args.cadastro)
    ativos, inativos, apelidos = ru.ler_cadastro(args.cadastro)

    ranking = calcular(dia, entradas, ativos, inativos, apelidos, config, feriados)
    saida = Path(args.saida) / f"ranking_vendas_{dia:%Y-%m-%d}.xlsx"
    gerar_planilha(ranking, config, saida, hoje)
    if usados is not None:
        usados.extend(ranking.usados)

    # Resumo no terminal
    print(f"\nRANKING DE CONTRATOS PAGOS — {dia.replace(day=1):%d/%m} a {dia:%d/%m/%Y} "
          f"({ranking.dias_ate} de {ranking.dias_mes} dias úteis)")
    for nome, obs in ranking.arquivos:
        print(f"  - {nome}: {obs}")
    for rotulo, qtd in ranking.contagem.items():
        print(f"  {rotulo}: {qtd}")
    print("\nTop 3 vendedores:")
    for l in ranking.vendedores[:3]:
        print(f"  {l.posicao}º {l.nome}: {ru.moeda(l.portabilidade)} ({l.projetando:.0%})")
    t = ranking.total
    print(f"\nTotal geral: {ru.moeda(t['portabilidade'])} ({t['projetando']:.0%} do esperado) | "
          f"andamento {ru.moeda(t['andamento'])} | elegíveis ao bônus: {t['elegiveis']}")
    for aviso in ranking.avisos:
        print(f"AVISO: {aviso}")
    print(f"\nRanking: {saida}")
    return 0


if __name__ == "__main__":
    try:
        sys.exit(main())
    except PermissionError as erro:
        print(f"\nERRO: não consegui gravar {erro.filename}.\n"
              "Se esse arquivo estiver aberto no Excel, feche-o e rode de novo.")
        sys.exit(1)
    except (FileNotFoundError, ValueError) as erro:
        print(f"\nERRO: {erro}")
        sys.exit(1)
