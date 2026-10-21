"""
Relatório diário de produtividade da URA.

Junta as exportações do dia e gera saida/relatorio_produtividade_AAAA-MM-DD.xlsx
no mesmo formato da planilha Relatorio_Produtividade_URA_Automatico.xlsx.

De onde vem cada coluna:
  TMA e TMO ............. Argus > Desempenho Operadores (colunas TMA e TME)
  Leads atendidos/URA ... Argus > Desempenho Operadores > grupo RECEPTIVO - URA
  Custo total ........... Tarifação VENDITORE + 4NET: custo do operador
                          + rateio do custo dos robôs da URA, na proporção dos leads
  Vendas geral .......... Vanguard: soma do Valor Contrato (etapas Andamento, Pendente, Pago)
  Clientes efetivados ... Vanguard: CPFs únicos com etapa Andamento ou Pago
  Custo por CPF ......... custo total / clientes efetivados
  Produção/lead ......... vendas geral / leads atendidos

Ordem: operadores de quem mais vendeu para quem menos vendeu (empate: mais leads),
depois o total de cada equipe (também por vendas) e os totais gerais.
Nas equipes, TMA e TMO são a SOMA dos operadores, conforme o procedimento.

Os arquivos são reconhecidos pelas COLUNAS, então o nome pode ser qualquer um.
Se o desempenho e os leads do dia não estiverem em entrada/ e o token do Argus
estiver configurado, eles são baixados automaticamente pela API.

Como usar:
  python relatorio_ura.py                    -> relatório do último dia útil
  python relatorio_ura.py --data 17/10/2026  -> relatório de outro dia
  python relatorio_ura.py --sem-api          -> não baixa nada do Argus, usa só entrada/
  python relatorio_ura.py --arquivar         -> no final, move os arquivos usados para
                                                entrada/processados/AAAA-MM-DD/

Extras:
  - Cada relatório é guardado em historico/relatorio_ura.csv. A partir do segundo dia,
    aparecem colunas comparando leads, vendas e custo com o relatório anterior.
  - Metas em config/metas.xlsx pintam as células dos operadores de verde, amarelo ou vermelho.
"""

import argparse
import csv
import shutil
import sys
from dataclasses import dataclass, field
from datetime import date, datetime, time, timedelta
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

from calendario import dia_util_anterior, ler_feriados
from esteira import nome_do_agente, remover_propostas_repetidas
from leitura import (converter_data, converter_valor_br, limpar_cabecalho, ler_tabela,
                     normalizar_nome, tempo_para_segundos)

# ---------------------------------------------------------------------------
# Configurações (mude aqui se a regra mudar)
# ---------------------------------------------------------------------------

PASTA_PROJETO = Path(__file__).resolve().parent
PASTA_ENTRADA = PASTA_PROJETO / "entrada"
PASTA_SAIDA = PASTA_PROJETO / "saida"
ARQUIVO_CADASTRO = PASTA_PROJETO / "config" / "operadores.xlsx"
ARQUIVO_METAS = PASTA_PROJETO / "config" / "metas.xlsx"
ARQUIVO_HISTORICO = PASTA_PROJETO / "historico" / "relatorio_ura.csv"

ETAPAS_VENDA = {"ANDAMENTO", "PENDENTE", "PAGO"}     # entram em Vendas geral
ETAPAS_CLIENTE = {"ANDAMENTO", "PAGO"}               # entram em Clientes efetivados
COLUNA_VALOR_VENDA = "VALOR CONTRATO"                # ou "VL. LIBERADA"

PREFIXOS_ROBO = ("OPER VIRT", "OPERADOR VIRT")       # linhas dos robôs da URA na tarifação
TEXTO_SEM_DERIVACAO = "SEM DERIVACAO"                # linha "<LIGAÇÕES SEM DERIVAÇÃO P/ OPERADOR>"

MOSTRAR_TODOS = False  # False: só operadores com leads, custo, vendas ou clientes

# Cores da empresa (as mesmas da planilha modelo)
AZUL = "0B2A5B"
AMARELO = "FFE600"
BRANCO = "FFFFFF"
CINZA_EQUIPE = "EDEDED"
CINZA_TOTAL = "D9D9D9"
ZEBRA = "F4F7FB"
AMARELO_CLARO = "FFF2CC"
MARROM = "7F6000"
# Semáforo das metas (fundo, fonte), nas cores padrão do Excel
SEMAFORO = {
    "verde": ("C6EFCE", "006100"),
    "amarelo": ("FFEB9C", "9C5700"),
    "vermelho": ("FFC7CE", "9C0006"),
}

# Colunas que identificam cada exportação (nomes padronizados, sem acento)
TIPOS_DE_ARQUIVO = {
    "esteira": {"CPF", "ETAPA", "AGENTE", "DATA/HORA"},
    "operadores": {"NOME", "SUPERVISOR(ES)", "STATUS"},
    "tarifacao": {"DESCRICAO", "VALOR R$", "MINUTOS"},
    "desempenho": {"DATA", "OPERADOR", "TME", "TMA"},
    "leads": {"DATA", "OPERADOR", "LEADS. ATEND."},
}
NOMES_DOS_TIPOS = {
    "esteira": "Esteira Vanguard", "operadores": "Operadores Argus", "tarifacao": "Tarifação Argus",
    "desempenho": "Desempenho Argus (TMA/TME)", "leads": "Leads URA Argus",
}


# ---------------------------------------------------------------------------
# Estruturas de dados
# ---------------------------------------------------------------------------

@dataclass
class Operador:
    """Um operador do cadastro e os números dele no dia."""
    nome: str
    equipe: str
    ordem: int                  # posição no cadastro (desempate final)
    tma: int = 0                # segundos
    tmo: int = 0                # segundos (coluna TME do Argus)
    leads: int = 0
    custo_direto: float = 0.0
    rateio_robos: float = 0.0
    vendas: float = 0.0
    cpfs: set = field(default_factory=set)
    tem_desempenho: bool = False

    @property
    def custo(self):
        return self.custo_direto + self.rateio_robos

    @property
    def clientes(self):
        return len(self.cpfs)

    def tem_atividade(self):
        return self.leads > 0 or self.custo > 0 or self.vendas > 0 or self.clientes > 0


@dataclass
class Resultado:
    """Tudo o que o relatório precisa, mais as informações de conferência."""
    data: date
    operadores: dict                 # {NOME NORMALIZADO: Operador}
    sem_derivacao: float = 0.0
    custo_robos: float = 0.0
    custo_importado: float = 0.0     # tarifação sem linhas de total
    custo_ignorado: float = 0.0      # tarifação de quem está fora do cadastro
    vendas_ignoradas: float = 0.0    # vendas de agentes fora do cadastro
    vendas_no_dia: int = 0           # linhas da esteira com a data do relatório
    sem_vinculo: dict = field(default_factory=dict)   # {nome: {origens}}
    ex_colaboradores: dict = field(default_factory=dict)
    avisos: list = field(default_factory=list)
    arquivos: list = field(default_factory=list)      # [(arquivo, tipo, observação)]
    usados: list = field(default_factory=list)        # caminhos dos arquivos usados (para arquivar)


# ---------------------------------------------------------------------------
# Leitura das entradas
# ---------------------------------------------------------------------------

def tipo_do_arquivo(cabecalho):
    """Descobre que exportação é pelo conjunto de colunas (ou None)."""
    colunas = {limpar_cabecalho(c) for c in cabecalho}
    for tipo, obrigatorias in TIPOS_DE_ARQUIVO.items():
        if obrigatorias <= colunas:
            return tipo
    return None


def como_dicionarios(cabecalho, linhas):
    """Transforma as linhas em dicionários {COLUNA PADRONIZADA: valor}."""
    colunas = [limpar_cabecalho(c) for c in cabecalho]
    return [dict(zip(colunas, linha)) for linha in linhas]


def ler_entradas(pasta):
    """
    Lê todos os arquivos da pasta e separa por tipo.

    Devolve ({tipo: [(arquivo, [linhas como dicionário])]}, [arquivos não reconhecidos]).
    """
    por_tipo = {tipo: [] for tipo in TIPOS_DE_ARQUIVO}
    desconhecidos = []
    for arquivo in sorted(Path(pasta).glob("*")):
        if not arquivo.is_file() or arquivo.suffix.lower() not in (".csv", ".xlsx", ".xls"):
            continue
        if arquivo.name.startswith("~$"):  # arquivo temporário do Excel aberto
            continue
        cabecalho, linhas = ler_tabela(arquivo)
        tipo = tipo_do_arquivo(cabecalho)
        if tipo:
            por_tipo[tipo].append((arquivo, como_dicionarios(cabecalho, linhas)))
        else:
            desconhecidos.append(arquivo)
    return por_tipo, desconhecidos


# ---------------------------------------------------------------------------
# Cadastro de operadores e apelidos (config/operadores.xlsx)
# ---------------------------------------------------------------------------

def limpar_supervisor(texto):
    """'(1) FULANO DE TAL, (2) OUTRO' -> 'FULANO DE TAL'."""
    primeiro = str(texto or "").split(",")[0].strip()
    if primeiro.startswith("(") and ")" in primeiro:
        primeiro = primeiro.split(")", 1)[1].strip()
    return primeiro


def nome_da_equipe(supervisor):
    """A equipe é o primeiro nome do supervisor: 'FULANO DE TAL' -> 'Fulano'."""
    partes = limpar_supervisor(supervisor).split()
    return partes[0].capitalize() if partes else ""


def atualizar_cadastro(linhas_exportacao, arquivo=ARQUIVO_CADASTRO):
    """
    Grava a aba Operadores de config/operadores.xlsx a partir da exportação do Argus.

    A aba Apelidos (preenchida por você) é mantida como está.
    Se o arquivo não existir, ele é criado com uma aba Apelidos vazia.
    """
    arquivo = Path(arquivo)
    if arquivo.exists():
        wb = openpyxl.load_workbook(arquivo)
        if "Operadores" in wb.sheetnames:
            del wb["Operadores"]
    else:
        arquivo.parent.mkdir(parents=True, exist_ok=True)
        wb = openpyxl.Workbook()
        apelidos = wb.active
        apelidos.title = "Apelidos"
        apelidos.append(["Nome no sistema", "Nome oficial"])
        apelidos.append(["(exemplo) JOANA SOUSA", "(exemplo) JOANNA SOUSA"])
        apelidos.column_dimensions["A"].width = 45
        apelidos.column_dimensions["B"].width = 45

    aba = wb.create_sheet("Operadores", 0)
    aba.append(["Nome", "Supervisor", "Equipe", "Status"])
    for linha in linhas_exportacao:
        nome = " ".join(str(linha.get("NOME", "")).split())
        if not nome:
            continue
        supervisor = limpar_supervisor(linha.get("SUPERVISOR(ES)"))
        status = normalizar_nome(linha.get("STATUS")) or "ATIVO"
        aba.append([nome, supervisor, nome_da_equipe(supervisor), status])
    for coluna, largura in zip("ABCD", (45, 40, 15, 10)):
        aba.column_dimensions[coluna].width = largura
    wb.save(arquivo)


def ler_cadastro(arquivo=ARQUIVO_CADASTRO):
    """
    Lê config/operadores.xlsx.

    Devolve (ativos, inativos, apelidos):
      ativos   = {NOME NORMALIZADO: (nome, equipe)} na ordem do cadastro
      inativos = {NOME NORMALIZADO} de quem tem Status diferente de ATIVO
      apelidos = {NOME NO SISTEMA NORMALIZADO: NOME OFICIAL NORMALIZADO}
    """
    arquivo = Path(arquivo)
    if not arquivo.exists():
        raise FileNotFoundError(
            f"Não encontrei o cadastro {arquivo}.\n"
            "Coloque a exportação de OPERADORES do Argus em entrada/ (ou configure o token "
            "da API) e rode de novo: o cadastro é criado sozinho."
        )
    wb = openpyxl.load_workbook(arquivo, read_only=True)
    ativos, inativos, apelidos = {}, set(), {}
    for nome, _supervisor, equipe, status in list(wb["Operadores"].iter_rows(min_row=2, values_only=True)):
        if not nome:
            continue
        chave = normalizar_nome(nome)
        if normalizar_nome(status) == "ATIVO":
            ativos.setdefault(chave, (str(nome).strip(), equipe or "SEM EQUIPE"))
        else:
            inativos.add(chave)
    if "Apelidos" in wb.sheetnames:
        for linha in wb["Apelidos"].iter_rows(min_row=2, values_only=True):
            no_sistema, oficial = (list(linha) + [None, None])[:2]
            if no_sistema and oficial and not str(no_sistema).startswith("(exemplo)"):
                apelidos[normalizar_nome(no_sistema)] = normalizar_nome(oficial)
    return ativos, inativos - set(ativos), apelidos


# ---------------------------------------------------------------------------
# Cálculos
# ---------------------------------------------------------------------------

class Vinculador:
    """
    Liga o nome que aparece em cada sistema ao operador do cadastro.

    Só usa o nome padronizado (sem acentos/maiúsculas) e a aba Apelidos.
    NÃO faz correspondência por semelhança: nome diferente vira aviso.
    """

    def __init__(self, resultado, inativos, apelidos):
        self.resultado = resultado
        self.inativos = inativos
        self.apelidos = apelidos

    def operador(self, nome, origem):
        """Devolve o Operador do cadastro, ou None (e registra o motivo)."""
        chave = normalizar_nome(nome)
        if not chave:
            return None
        chave = self.apelidos.get(chave, chave)
        if chave in self.resultado.operadores:
            return self.resultado.operadores[chave]
        lista = self.resultado.ex_colaboradores if chave in self.inativos else self.resultado.sem_vinculo
        lista.setdefault(" ".join(str(nome).split()), set()).add(origem)
        return None


def linhas_do_dia(arquivos, dia, resultado, nome_tipo):
    """Junta as linhas com Data = dia; avisa sobre arquivos de outros dias."""
    escolhidas = []
    for arquivo, linhas in arquivos:
        do_dia = [l for l in linhas if converter_data(l.get("DATA")) == dia]
        datas = sorted({converter_data(l.get("DATA")) for l in linhas} - {None})
        if do_dia:
            escolhidas += do_dia
            resultado.usados.append(arquivo)
            resultado.arquivos.append((arquivo.name, nome_tipo, f"{len(do_dia)} linha(s) de {dia:%d/%m/%Y}"))
        else:
            outras = ", ".join(f"{d:%d/%m/%Y}" for d in datas) or "sem data"
            resultado.arquivos.append((arquivo.name, nome_tipo, f"IGNORADO: é de {outras}"))
    return escolhidas


def calcular(dia, entradas, ativos, inativos, apelidos):
    """Faz todas as contas do relatório e devolve um Resultado."""
    operadores = {chave: Operador(nome, equipe, i) for i, (chave, (nome, equipe)) in enumerate(ativos.items())}
    resultado = Resultado(data=dia, operadores=operadores)
    vincular = Vinculador(resultado, inativos, apelidos)

    # 1. TMA e TMO (desempenho do Argus)
    for linha in linhas_do_dia(entradas["desempenho"], dia, resultado, NOMES_DOS_TIPOS["desempenho"]):
        op = vincular.operador(linha.get("OPERADOR"), "Argus desempenho")
        if op:
            op.tma = tempo_para_segundos(linha.get("TMA"))
            op.tmo = tempo_para_segundos(linha.get("TME"))
            op.tem_desempenho = True

    # 2. Leads atendidos da URA
    for linha in linhas_do_dia(entradas["leads"], dia, resultado, NOMES_DOS_TIPOS["leads"]):
        op = vincular.operador(linha.get("OPERADOR"), "Argus leads URA")
        if op:
            op.leads += int(converter_valor_br(linha.get("LEADS. ATEND.")))

    # 3. Custos (tarifação VENDITORE + 4NET)
    for arquivo, linhas in entradas["tarifacao"]:
        resultado.arquivos.append((arquivo.name, NOMES_DOS_TIPOS["tarifacao"], f"{len(linhas)} linha(s)"))
        resultado.usados.append(arquivo)
        for linha in linhas:
            descricao = str(linha.get("DESCRICAO", "")).strip()
            valor = converter_valor_br(linha.get("VALOR R$"))
            if descricao.startswith("<"):
                # Linha de total; só a "SEM DERIVAÇÃO" entra no custo
                if TEXTO_SEM_DERIVACAO in normalizar_nome(descricao):
                    resultado.sem_derivacao += valor
                    resultado.custo_importado += valor
                continue
            if not descricao:
                continue
            resultado.custo_importado += valor
            if normalizar_nome(descricao).startswith(PREFIXOS_ROBO):
                resultado.custo_robos += valor
                continue
            op = vincular.operador(descricao, "Tarifação")
            if op:
                op.custo_direto += valor
            else:
                resultado.custo_ignorado += valor

    # 4. Rateio do custo dos robôs na proporção dos leads da URA
    total_leads = sum(op.leads for op in operadores.values())
    if total_leads > 0:
        for op in operadores.values():
            op.rateio_robos = resultado.custo_robos * op.leads / total_leads
    elif resultado.custo_robos > 0:
        resultado.avisos.append(
            f"Há R$ {resultado.custo_robos:.2f} de custo dos robôs da URA, mas nenhum lead "
            "atendido para fazer o rateio. Esse custo ficou fora do relatório.")

    # 5. Vendas e clientes efetivados (esteira do Vanguard)
    # Podem vir duas esteiras (data de andamento e data de pagamento): junta as duas
    # e conta cada proposta uma vez só.
    todas = []
    for arquivo, linhas in entradas["esteira"]:
        resultado.usados.append(arquivo)
        resultado.arquivos.append((arquivo.name, NOMES_DOS_TIPOS["esteira"], f"{len(linhas)} proposta(s)"))
        todas += linhas
    propostas, repetidas = remover_propostas_repetidas(todas)
    if repetidas:
        resultado.avisos.append(f"{repetidas} proposta(s) estavam em mais de uma esteira e foram contadas uma vez só.")
    do_dia = [l for l in propostas if converter_data(l.get("DATA/HORA")) == dia]
    resultado.vendas_no_dia = len(do_dia)
    for linha in do_dia:
        etapa = normalizar_nome(linha.get("ETAPA"))
        if etapa not in ETAPAS_VENDA and etapa not in ETAPAS_CLIENTE:
            continue
        valor = converter_valor_br(linha.get(COLUNA_VALOR_VENDA))
        op = vincular.operador(nome_do_agente(linha.get("AGENTE")), "Vanguard")
        if op is None:
            if etapa in ETAPAS_VENDA:
                resultado.vendas_ignoradas += valor
            continue
        if etapa in ETAPAS_VENDA:
            op.vendas += valor
        if etapa in ETAPAS_CLIENTE and str(linha.get("CPF", "")).strip():
            op.cpfs.add(str(linha["CPF"]).strip())

    # Conferências
    if not entradas["desempenho"] or not any(op.tem_desempenho for op in operadores.values()):
        resultado.avisos.append(f"Não há desempenho (TMA/TME) de {dia:%d/%m/%Y} em entrada/.")
    if total_leads == 0:
        resultado.avisos.append(f"Não há leads da URA de {dia:%d/%m/%Y} em entrada/.")
    if not entradas["tarifacao"]:
        resultado.avisos.append("Não há tarifação (VENDITORE/4NET) em entrada/: o custo ficou zerado.")
    elif len(entradas["tarifacao"]) > 2:
        resultado.avisos.append(f"Há {len(entradas['tarifacao'])} tarifações em entrada/ (o normal são 2: "
                                "VENDITORE e 4NET). Todas foram somadas; apague as de outros dias.")
    if not entradas["esteira"]:
        resultado.avisos.append("Não há esteira do Vanguard em entrada/: vendas e clientes ficaram zerados.")
    elif resultado.vendas_no_dia == 0:
        resultado.avisos.append(f"A esteira não tem nenhuma proposta com data {dia:%d/%m/%Y}.")
    return resultado


# ---------------------------------------------------------------------------
# Montagem das linhas do relatório
# ---------------------------------------------------------------------------

def dividir(a, b):
    """Divisão que devolve 0 quando o divisor é 0 (igual ao SEERRO do Excel)."""
    return a / b if b else 0


def linhas_do_relatorio(resultado, mostrar_todos=MOSTRAR_TODOS):
    """
    Monta as linhas na ordem do relatório.

    Cada linha é um dicionário com 'tipo' (operador, separador, equipe, total_operadores,
    sem_derivacao, total_geral) e os valores das colunas.
    """
    visiveis = [op for op in resultado.operadores.values() if mostrar_todos or op.tem_atividade()]
    # Quem mais vendeu primeiro; empate: mais leads; depois a ordem do cadastro
    visiveis.sort(key=lambda op: (-op.vendas, -op.leads, op.ordem))

    def valores(tma, tmo, leads, custo, vendas, clientes):
        return {"tma": tma, "tmo": tmo, "leads": leads, "custo": custo, "vendas": vendas,
                "clientes": clientes, "custo_cpf": dividir(custo, clientes), "prod_lead": dividir(vendas, leads)}

    linhas = [{"tipo": "operador", "nome": op.nome,
               **valores(op.tma, op.tmo, op.leads, op.custo, op.vendas, op.clientes)} for op in visiveis]
    linhas.append({"tipo": "separador", "nome": ""})

    # Equipes: somam os operadores visíveis (TMA e TMO também são somados)
    equipes = {}
    for op in visiveis:
        equipes.setdefault(op.equipe, []).append(op)
    totais_equipe = []
    for equipe, membros in equipes.items():
        totais_equipe.append({"tipo": "equipe", "nome": f"TOTAL EQUIPE - {equipe}", **valores(
            sum(m.tma for m in membros), sum(m.tmo for m in membros), sum(m.leads for m in membros),
            sum(m.custo for m in membros), sum(m.vendas for m in membros), sum(m.clientes for m in membros))})
    totais_equipe.sort(key=lambda e: (-e["vendas"], e["nome"]))
    linhas += totais_equipe

    leads = sum(op.leads for op in visiveis)
    custo = sum(op.custo for op in visiveis)
    vendas = sum(op.vendas for op in visiveis)
    clientes = sum(op.clientes for op in visiveis)
    linhas.append({"tipo": "total_operadores", "nome": "TOTAL OPERADORES",
                   **valores(None, None, leads, custo, vendas, clientes)})
    linhas.append({"tipo": "sem_derivacao", "nome": "LIGAÇÕES SEM DERIVAÇÃO P/ OPERADOR",
                   "custo": resultado.sem_derivacao})
    linhas.append({"tipo": "total_geral", "nome": "TOTAL GERAL (c/ sem derivação)",
                   **valores(None, None, leads, custo + resultado.sem_derivacao, vendas, clientes)})
    return linhas


# ---------------------------------------------------------------------------
# Histórico e comparação com o relatório anterior
# ---------------------------------------------------------------------------

COLUNAS_HISTORICO = ["Data", "Tipo", "Nome", "TMA (s)", "TMO (s)", "Leads", "Custo", "Vendas", "Clientes"]
TIPOS_NO_HISTORICO = ("operador", "equipe", "total_operadores", "total_geral")


def numero_br(valor, casas=2):
    """1234.5 -> '1234,50' (vírgula decimal, para o Excel brasileiro abrir certo)."""
    return f"{valor:.{casas}f}".replace(".", ",")


def ler_historico(arquivo=ARQUIVO_HISTORICO):
    """Lê historico/relatorio_ura.csv: lista de dicionários com os números já convertidos."""
    arquivo = Path(arquivo)
    if not arquivo.exists():
        return []
    with open(arquivo, encoding="utf-8-sig", newline="") as f:
        return [{
            "data": converter_data(l["Data"]), "tipo": l["Tipo"], "nome": l["Nome"],
            "tma": int(converter_valor_br(l["TMA (s)"])), "tmo": int(converter_valor_br(l["TMO (s)"])),
            "leads": int(converter_valor_br(l["Leads"])), "custo": converter_valor_br(l["Custo"]),
            "vendas": converter_valor_br(l["Vendas"]), "clientes": int(converter_valor_br(l["Clientes"])),
        } for l in csv.DictReader(f, delimiter=";")]


def salvar_no_historico(dia, linhas, arquivo=ARQUIVO_HISTORICO):
    """
    Acrescenta os números do dia ao histórico (uma linha por operador, equipe e total).

    Se o dia já estiver lá (relatório rodado de novo), as linhas antigas são substituídas.
    """
    arquivo = Path(arquivo)
    anteriores = [r for r in ler_historico(arquivo) if r["data"] != dia]
    novos = [{"data": dia, "tipo": l["tipo"], "nome": l["nome"], "tma": l.get("tma") or 0,
              "tmo": l.get("tmo") or 0, "leads": l.get("leads") or 0, "custo": l.get("custo") or 0,
              "vendas": l.get("vendas") or 0, "clientes": l.get("clientes") or 0}
             for l in linhas if l["tipo"] in TIPOS_NO_HISTORICO]
    arquivo.parent.mkdir(parents=True, exist_ok=True)
    with open(arquivo, "w", encoding="utf-8-sig", newline="") as f:
        escritor = csv.writer(f, delimiter=";")
        escritor.writerow(COLUNAS_HISTORICO)
        for r in sorted(anteriores + novos, key=lambda r: r["data"]):
            escritor.writerow([f"{r['data']:%d/%m/%Y}", r["tipo"], r["nome"], r["tma"], r["tmo"],
                               r["leads"], numero_br(r["custo"]), numero_br(r["vendas"]), r["clientes"]])


def comparar_com_anterior(linhas, historico, dia, esperado=None, avisos=None):
    """
    Acrescenta em cada linha a diferença em relação ao relatório do dia útil anterior:
    d_leads, d_vendas e d_custo. Devolve a data comparada (ou None se não houver).

    `esperado` é o dia útil anterior ao dia do relatório (ex.: relatório de segunda,
    19/10 -> compara com sexta, 16/10). Se esse dia não estiver no histórico, compara
    com o relatório mais recente antes de `dia` e registra um aviso em `avisos`.

    Quem não aparecia no relatório anterior fica sem comparação.
    """
    datas = sorted({r["data"] for r in historico if r["data"] and r["data"] < dia})
    if not datas:
        return None
    if esperado in datas:
        anterior = esperado
    else:
        anterior = datas[-1]
        if esperado and avisos is not None:
            avisos.append(f"O relatório de {esperado:%d/%m/%Y} (dia útil anterior) não está no histórico: "
                          f"a comparação foi feita com o de {anterior:%d/%m/%Y}. Para gerá-lo: "
                          f"python relatorio_ura.py --data {esperado:%d/%m/%Y}")
    do_anterior = {(r["tipo"], normalizar_nome(r["nome"])): r for r in historico if r["data"] == anterior}
    for linha in linhas:
        antes = do_anterior.get((linha["tipo"], normalizar_nome(linha["nome"])))
        if antes:
            linha["d_leads"] = (linha.get("leads") or 0) - antes["leads"]
            linha["d_vendas"] = (linha.get("vendas") or 0) - antes["vendas"]
            linha["d_custo"] = (linha.get("custo") or 0) - antes["custo"]
    return anterior


# ---------------------------------------------------------------------------
# Metas (semáforo verde / amarelo / vermelho)
# ---------------------------------------------------------------------------

# Nome do indicador na planilha de metas -> chave da coluna no relatório
INDICADORES = {
    "TMA": "tma", "TMO": "tmo", "LEADS ATENDIDOS/URA": "leads", "CUSTO TOTAL": "custo",
    "VENDAS GERAL": "vendas", "CLIENTE EFETIVADOS": "clientes", "CUSTO POR CPF": "custo_cpf",
    "PRODUCAO/LEAD": "prod_lead",
}


def criar_modelo_metas(arquivo=ARQUIVO_METAS):
    """Cria config/metas.xlsx com os indicadores e os limites em branco, se ainda não existir."""
    arquivo = Path(arquivo)
    if arquivo.exists():
        return False
    arquivo.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Metas"
    ws.append(["Indicador", "Melhor quando", "Verde", "Vermelho"])
    for celula in ws[1]:
        preencher(celula, AZUL, BRANCO, True)
    for indicador, sentido in [("TMA", "menor"), ("TMO", "menor"), ("LEADS ATENDIDOS/URA", "maior"),
                               ("CUSTO TOTAL", "menor"), ("VENDAS GERAL", "maior"),
                               ("CLIENTE EFETIVADOS", "maior"), ("CUSTO POR CPF", "menor"),
                               ("PRODUÇÃO/LEAD", "maior")]:
        ws.append([indicador, sentido, None, None])
    instrucoes = [
        "COMO PREENCHER",
        "Deixe Verde e Vermelho em branco para não pintar aquele indicador.",
        "Melhor quando = maior: verde se o valor for MAIOR OU IGUAL ao Verde;",
        "   vermelho se for MENOR OU IGUAL ao Vermelho; amarelo no meio.",
        "Melhor quando = menor: verde se o valor for MENOR OU IGUAL ao Verde;",
        "   vermelho se for MAIOR OU IGUAL ao Vermelho; amarelo no meio.",
        "TMA e TMO: escreva o tempo como 00:01:30 (ou em segundos: 90).",
        "Valores em reais: só o número (ex.: 15000 ou 15000,50).",
        "As cores valem para as linhas dos operadores.",
    ]
    for i, texto in enumerate(instrucoes, start=1):
        ws.cell(i, 6, texto).font = Font(bold=(i == 1), color=AZUL if i == 1 else None)
    for coluna, largura in zip("ABCDEF", (24, 16, 12, 12, 3, 70)):
        ws.column_dimensions[coluna].width = largura
    wb.save(arquivo)
    return True


def valor_da_meta(valor, chave):
    """Converte o limite digitado na planilha de metas em número (segundos para TMA/TMO)."""
    if valor is None or str(valor).strip() == "":
        return None
    if chave in ("tma", "tmo"):
        if isinstance(valor, timedelta):
            return valor.total_seconds()
        if isinstance(valor, time):
            return valor.hour * 3600 + valor.minute * 60 + valor.second
        if isinstance(valor, datetime):
            return valor.hour * 3600 + valor.minute * 60 + valor.second
        if isinstance(valor, float) and valor < 1:  # o Excel guardou como fração do dia
            return round(valor * 86400)
        if ":" in str(valor):
            return tempo_para_segundos(valor)
    return converter_valor_br(valor)


def ler_metas(arquivo=ARQUIVO_METAS):
    """
    Lê config/metas.xlsx: {chave: (sentido, verde, vermelho)}.

    Só entram os indicadores com pelo menos um limite preenchido.
    """
    arquivo = Path(arquivo)
    if not arquivo.exists():
        return {}
    metas = {}
    wb = openpyxl.load_workbook(arquivo, read_only=True, data_only=True)
    for linha in wb.worksheets[0].iter_rows(min_row=2, max_col=4, values_only=True):
        indicador, sentido, verde, vermelho = (list(linha) + [None] * 4)[:4]
        chave = INDICADORES.get(normalizar_nome(indicador))
        if not chave:
            continue
        verde, vermelho = valor_da_meta(verde, chave), valor_da_meta(vermelho, chave)
        if verde is None and vermelho is None:
            continue
        sentido = "menor" if normalizar_nome(sentido) == "MENOR" else "maior"
        metas[chave] = (sentido, verde, vermelho)
    return metas


def cor_da_meta(valor, regra):
    """'verde', 'amarelo' ou 'vermelho' para o valor, conforme a regra (sentido, verde, vermelho)."""
    if valor is None or regra is None:
        return None
    sentido, verde, vermelho = regra
    if sentido == "maior":
        if verde is not None and valor >= verde:
            return "verde"
        if vermelho is not None and valor <= vermelho:
            return "vermelho"
    else:
        if verde is not None and valor <= verde:
            return "verde"
        if vermelho is not None and valor >= vermelho:
            return "vermelho"
    return "amarelo"


# ---------------------------------------------------------------------------
# Arquivamento da pasta entrada/
# ---------------------------------------------------------------------------

def arquivar(arquivos, pasta_entrada, dia):
    """
    Move os arquivos usados para entrada/processados/AAAA-MM-DD/.

    Assim a próxima execução não soma, por engano, a tarifação de hoje.
    O conteúdo dos arquivos não é alterado. Devolve a pasta de destino.
    """
    destino = Path(pasta_entrada) / "processados" / f"{dia:%Y-%m-%d}"
    destino.mkdir(parents=True, exist_ok=True)
    for arquivo in dict.fromkeys(Path(a) for a in arquivos):  # sem repetidos, na ordem
        if not arquivo.exists():
            continue
        alvo = destino / arquivo.name
        n = 2
        while alvo.exists():  # não sobrescreve: acrescenta _2, _3...
            alvo = destino / f"{arquivo.stem}_{n}{arquivo.suffix}"
            n += 1
        shutil.move(str(arquivo), str(alvo))
    return destino


# ---------------------------------------------------------------------------
# Planilha de saída
# ---------------------------------------------------------------------------

COLUNAS = [  # (título, chave, formato)
    ("TMA", "tma", "[h]:mm:ss"),
    ("TMO", "tmo", "[h]:mm:ss"),
    ("LEADS ATENDIDOS/URA", "leads", "0"),
    ("CUSTO TOTAL", "custo", '"R$ "#,##0.00'),
    ("VENDAS GERAL", "vendas", '"R$ "#,##0.00'),
    ("CLIENTE EFETIVADOS", "clientes", "0"),
    ("CUSTO POR CPF", "custo_cpf", '"R$ "#,##0.00'),
    ("PRODUÇÃO/LEAD", "prod_lead", '"R$ "#,##0.00'),
]
# Colunas extras de comparação (só aparecem se houver relatório anterior no histórico).
# ▲ = subiu, ▼ = desceu. A cor diz se foi bom (verde) ou ruim (vermelho):
# leads e vendas subindo é bom; custo subindo é ruim.
FORMATO_DIF_LEADS = '"▲ "#,##0;"▼ "#,##0;"="'
FORMATO_DIF_REAIS = '"▲ R$ "#,##0.00;"▼ R$ "#,##0.00;"="'
SUBIR_E_BOM = {"d_leads": True, "d_vendas": True, "d_custo": False}
VERDE_TEXTO = "006100"
VERMELHO_TEXTO = "C00000"


def colunas_de_comparacao(data_anterior):
    rotulo = f"{data_anterior:%d/%m}"
    return [
        (f"LEADS vs {rotulo}", "d_leads", FORMATO_DIF_LEADS),
        (f"VENDAS vs {rotulo}", "d_vendas", FORMATO_DIF_REAIS),
        (f"CUSTO vs {rotulo}", "d_custo", FORMATO_DIF_REAIS),
    ]


ESTILOS = {  # tipo de linha: (cor de fundo, cor da fonte, negrito)
    "operador": (None, None, False),
    "separador": (AMARELO, None, False),
    "equipe": (CINZA_EQUIPE, None, True),
    "total_operadores": (CINZA_TOTAL, AZUL, True),
    "sem_derivacao": (AMARELO_CLARO, MARROM, False),
    "total_geral": (AZUL, BRANCO, True),
}


def preencher(celula, cor_fundo=None, cor_fonte=None, negrito=False):
    if cor_fundo:
        celula.fill = PatternFill("solid", fgColor=cor_fundo)
    celula.font = Font(bold=negrito, color=cor_fonte)


def gerar_planilha(resultado, linhas, caminho, metas=None, data_anterior=None):
    """
    Cria o .xlsx do relatório (aba Relatório) e a aba Conferência.

    metas: {chave: regra} para pintar as células dos operadores (semáforo).
    data_anterior: se informada, acrescenta as colunas de comparação.
    """
    metas = metas or {}
    colunas = COLUNAS + (colunas_de_comparacao(data_anterior) if data_anterior else [])
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Relatório"
    ultima_coluna = 3 + len(colunas)  # A = faixa azul, B = nomes, C.. = valores, última = faixa azul

    # Título e cabeçalho
    ws.merge_cells(start_row=2, start_column=2, end_row=2, end_column=ultima_coluna - 1)
    ws.cell(2, 2, f"RELATORIO DE PRODUTIVIDADE URA {resultado.data:%d/%m/%Y}")
    ws.cell(2, 2).font = Font(bold=True, size=14, color=AZUL)
    ws.cell(2, 2).alignment = Alignment(horizontal="center")
    for i, titulo in enumerate(["OPERADORES"] + [c[0] for c in colunas]):
        celula = ws.cell(3, 2 + i, titulo)
        preencher(celula, AZUL, BRANCO, True)
        celula.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # Linhas
    linha_excel = 4
    for n, linha in enumerate(linhas):
        fundo, fonte, negrito = ESTILOS[linha["tipo"]]
        if linha["tipo"] == "operador" and n % 2 == 1:
            fundo = ZEBRA
        ws.cell(linha_excel, 2, linha["nome"])
        for coluna in range(2, ultima_coluna):
            preencher(ws.cell(linha_excel, coluna), fundo, fonte, negrito)
        for i, (_titulo, chave, formato) in enumerate(colunas):
            valor = linha.get(chave)
            celula = ws.cell(linha_excel, 3 + i)
            if valor is not None:
                # Tempos vão como fração do dia, que é como o Excel guarda horas
                celula.value = valor / 86400 if chave in ("tma", "tmo") else valor
                celula.number_format = formato
            # Comparação: verde se melhorou, vermelho se piorou
            if chave in SUBIR_E_BOM and valor and linha["tipo"] != "total_geral":
                melhorou = (valor > 0) == SUBIR_E_BOM[chave]
                celula.font = Font(bold=negrito, color=VERDE_TEXTO if melhorou else VERMELHO_TEXTO)
            # Semáforo das metas, só nas linhas de operador
            cor = cor_da_meta(valor, metas.get(chave)) if linha["tipo"] == "operador" else None
            if cor:
                fundo_meta, fonte_meta = SEMAFORO[cor]
                celula.fill = PatternFill("solid", fgColor=fundo_meta)
                celula.font = Font(color=fonte_meta, bold=True)
        linha_excel += 1
    # Faixa azul de fechamento e laterais (como no modelo)
    for coluna in range(1, ultima_coluna + 1):
        preencher(ws.cell(linha_excel, coluna), AZUL)
    for linha_n in range(3, linha_excel + 1):
        preencher(ws.cell(linha_n, 1), AZUL)
        preencher(ws.cell(linha_n, ultima_coluna), AZUL)

    ws.column_dimensions["A"].width = 2
    ws.column_dimensions["B"].width = 44
    for i, (_titulo, chave, _formato) in enumerate(colunas):
        ws.column_dimensions[openpyxl.utils.get_column_letter(3 + i)].width = 19 if chave == "leads" else 15
    ws.column_dimensions[openpyxl.utils.get_column_letter(ultima_coluna)].width = 2
    ws.row_dimensions[3].height = 45
    ws.freeze_panes = "C4"
    # Impressão: paisagem, todas as colunas em uma página de largura
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    # --- Aba de conferência ---
    wc = wb.create_sheet("Conferência")
    wc.column_dimensions["A"].width = 55
    wc.column_dimensions["B"].width = 40
    wc.column_dimensions["C"].width = 40

    def titulo(texto):
        wc.append([])
        wc.append([texto])
        preencher(wc.cell(wc.max_row, 1), AZUL, BRANCO, True)

    wc.append([f"Conferência do relatório de {resultado.data:%d/%m/%Y}"])
    wc.cell(1, 1).font = Font(bold=True, size=13, color=AZUL)
    titulo("Arquivos lidos")
    for nome, tipo, obs in resultado.arquivos:
        wc.append([nome, tipo, obs])
    titulo("Custos")
    custo_relatorio = sum(op.custo for op in resultado.operadores.values()) + resultado.sem_derivacao
    for rotulo, valor in [
        ("Custo importado da tarifação (sem linhas de total)", resultado.custo_importado),
        ("Custo no relatório (operadores + sem derivação)", custo_relatorio),
        ("Custo de quem está fora do cadastro (ignorado)", resultado.custo_ignorado),
        ("Diferença (importado - relatório - ignorado)", resultado.custo_importado - custo_relatorio - resultado.custo_ignorado),
        ("Custo dos robôs da URA (rateado pelos leads)", resultado.custo_robos),
        ("Ligações sem derivação p/ operador", resultado.sem_derivacao),
        ("Vendas de agentes fora do cadastro (ignoradas)", resultado.vendas_ignoradas),
    ]:
        wc.append([rotulo, round(valor, 2)])
        wc.cell(wc.max_row, 2).number_format = '"R$ "#,##0.00'
    titulo("Nomes sem vínculo (cadastre em Apelidos se for grafia diferente)")
    for nome, origens in sorted(resultado.sem_vinculo.items()):
        wc.append([nome, ", ".join(sorted(origens))])
    titulo("Ex-colaboradores ignorados (inativos no Argus)")
    for nome, origens in sorted(resultado.ex_colaboradores.items()):
        wc.append([nome, ", ".join(sorted(origens))])
    titulo("Metas aplicadas (config/metas.xlsx)")
    if not metas:
        wc.append(["Nenhuma meta preenchida."])
    nomes_das_colunas = {chave: titulo_coluna for titulo_coluna, chave, _ in COLUNAS}
    for chave, (sentido, verde, vermelho) in metas.items():
        wc.append([nomes_das_colunas[chave], f"melhor quando {sentido}",
                   f"verde: {verde if verde is not None else '-'} | vermelho: {vermelho if vermelho is not None else '-'}"])
    titulo("Avisos")
    for aviso in resultado.avisos:
        wc.append([aviso])

    caminho = Path(caminho)
    caminho.parent.mkdir(parents=True, exist_ok=True)
    wb.save(caminho)


# ---------------------------------------------------------------------------
# Programa principal
# ---------------------------------------------------------------------------

def moeda(valor):
    """1234.5 -> 'R$ 1.234,50'"""
    return f"R$ {valor:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def sinal_moeda(valor):
    """Diferença em reais com sinal: 150 -> '+R$ 150,00'; -20 -> '-R$ 20,00'."""
    return ("+" if valor >= 0 else "-") + moeda(abs(valor))


def baixar_do_argus_se_preciso(dia, entradas, pasta_entrada):
    """
    Se faltar desempenho ou leads do dia, tenta baixar pela API do Argus.

    Devolve True se baixou (aí as entradas precisam ser lidas de novo).
    """
    def tem_do_dia(tipo):
        return any(converter_data(l.get("DATA")) == dia for _, linhas in entradas[tipo] for l in linhas)

    if tem_do_dia("desempenho") and tem_do_dia("leads"):
        return False
    import argus_api  # só é carregado quando precisa

    try:
        token = argus_api.ler_token(argus_api.ARQUIVO_TOKEN)
    except argus_api.ErroArgus:
        print("Sem token do Argus configurado: usando só os arquivos de entrada/.")
        return False
    print(f"Desempenho/leads de {dia:%d/%m/%Y} não estão em entrada/: baixando do Argus...")
    try:
        argus_api.baixar(dia, token, pasta_entrada)
    except argus_api.ErroArgus as erro:
        print(f"AVISO: não consegui baixar do Argus: {erro}")
        return False
    return True


def main(argumentos=None, usados=None):
    """
    `usados`: se for uma lista, recebe os arquivos de entrada/ usados no relatório
    (o benzaiten.py usa isso para arquivar tudo no final da rotina).
    """
    parser = argparse.ArgumentParser(description="Relatório diário de produtividade da URA.")
    parser.add_argument("--data", help="dia do relatório dd/mm/aaaa (padrão: último dia útil)")
    parser.add_argument("--sem-api", action="store_true", help="não baixa nada do Argus")
    parser.add_argument("--arquivar", action="store_true",
                        help="no final, move os arquivos usados para entrada/processados/AAAA-MM-DD/")
    parser.add_argument("--entrada", default=str(PASTA_ENTRADA), help=argparse.SUPPRESS)
    parser.add_argument("--saida", default=str(PASTA_SAIDA), help=argparse.SUPPRESS)
    parser.add_argument("--cadastro", default=str(ARQUIVO_CADASTRO), help=argparse.SUPPRESS)
    parser.add_argument("--historico", default=str(ARQUIVO_HISTORICO), help=argparse.SUPPRESS)
    parser.add_argument("--metas", default=str(ARQUIVO_METAS), help=argparse.SUPPRESS)
    args = parser.parse_args(argumentos)

    if args.data:
        try:
            dia = datetime.strptime(args.data, "%d/%m/%Y").date()
        except ValueError:
            raise ValueError(f"Data inválida: {args.data}. Use o formato dd/mm/aaaa, ex.: 17/10/2026")
    else:
        dia = dia_util_anterior(date.today(), ler_feriados())

    entrada = Path(args.entrada)
    if not entrada.exists():
        raise FileNotFoundError(f"A pasta {entrada} não existe. Crie-a e coloque as exportações do dia.")
    entradas, desconhecidos = ler_entradas(entrada)
    if not args.sem_api and baixar_do_argus_se_preciso(dia, entradas, entrada):
        entradas, desconhecidos = ler_entradas(entrada)

    # Cadastro: atualizado pela exportação de operadores mais recente, se houver
    if entradas["operadores"]:
        arquivo, linhas = max(entradas["operadores"], key=lambda par: par[0].stat().st_mtime)
        atualizar_cadastro(linhas, args.cadastro)
        print(f"Cadastro atualizado a partir de {arquivo.name}.")
    ativos, inativos, apelidos = ler_cadastro(args.cadastro)

    resultado = calcular(dia, entradas, ativos, inativos, apelidos)
    for arquivo, _ in entradas["operadores"]:
        resultado.arquivos.append((arquivo.name, NOMES_DOS_TIPOS["operadores"], "cadastro"))
        resultado.usados.append(arquivo)
    for arquivo in desconhecidos:
        resultado.arquivos.append((arquivo.name, "não reconhecido", "IGNORADO"))
    linhas = linhas_do_relatorio(resultado)

    # Comparação com o relatório anterior e metas
    # O relatório de um dia é comparado com o do dia útil anterior a ele
    # (ex.: feito em 20/10 -> relatório de 19/10 -> comparado com 16/10, a sexta)
    feriados = ler_feriados()
    data_anterior = comparar_com_anterior(linhas, ler_historico(args.historico), dia,
                                          dia_util_anterior(dia, feriados), resultado.avisos)
    if criar_modelo_metas(args.metas):
        print(f"Criado {args.metas}: preencha os limites para ativar o semáforo de metas.")
    metas = ler_metas(args.metas)

    saida = Path(args.saida) / f"relatorio_produtividade_{dia:%Y-%m-%d}.xlsx"
    gerar_planilha(resultado, linhas, saida, metas, data_anterior)
    salvar_no_historico(dia, linhas, args.historico)

    # Resumo no terminal
    print(f"\nRELATÓRIO DE PRODUTIVIDADE URA — {dia:%d/%m/%Y} (gerado em {date.today():%d/%m/%Y})")
    print("\nArquivos lidos:")
    for nome, tipo, obs in resultado.arquivos:
        print(f"  - {nome}: {tipo} ({obs})")
    total = next(l for l in linhas if l["tipo"] == "total_geral")
    total_op = next(l for l in linhas if l["tipo"] == "total_operadores")
    print(f"\nOperadores no relatório: {sum(1 for l in linhas if l['tipo'] == 'operador')} "
          f"(de {len(ativos)} ativos no cadastro)")
    print(f"Leads URA: {total['leads']} | Vendas: {moeda(total['vendas'])} | Clientes efetivados: {total['clientes']}")
    print(f"Custo operadores: {moeda(total_op['custo'])} | Sem derivação: {moeda(resultado.sem_derivacao)} "
          f"| Custo total: {moeda(total['custo'])}")
    if resultado.ex_colaboradores:
        print(f"\nEx-colaboradores ignorados: {', '.join(sorted(resultado.ex_colaboradores))}")
    if resultado.sem_vinculo:
        print("\nNOMES SEM VÍNCULO (ignorados; se for grafia diferente, cadastre na aba Apelidos):")
        for nome, origens in sorted(resultado.sem_vinculo.items()):
            print(f"  - {nome} ({', '.join(sorted(origens))})")
    if resultado.custo_ignorado or resultado.vendas_ignoradas:
        print(f"Ignorado por estar fora do cadastro: custo {moeda(resultado.custo_ignorado)}, "
              f"vendas {moeda(resultado.vendas_ignoradas)}")
    for aviso in resultado.avisos:
        print(f"AVISO: {aviso}")

    if data_anterior:
        print(f"\nComparado com {data_anterior:%d/%m/%Y}: "
              f"leads {total.get('d_leads', 0):+d} | vendas {sinal_moeda(total.get('d_vendas', 0))} "
              f"| custo {sinal_moeda(total.get('d_custo', 0))}")
    else:
        print("\nSem relatório anterior no histórico: a comparação aparece a partir do próximo.")
    if metas:
        contagem = {"verde": 0, "amarelo": 0, "vermelho": 0}
        for linha in linhas:
            if linha["tipo"] == "operador":
                for chave, regra in metas.items():
                    cor = cor_da_meta(linha.get(chave), regra)
                    if cor:
                        contagem[cor] += 1
        print(f"Metas: {contagem['verde']} verde(s), {contagem['amarelo']} amarela(s), "
              f"{contagem['vermelho']} vermelha(s)")

    print(f"\nRelatório: {saida}")
    print(f"Histórico: {args.historico}")
    if usados is not None:
        usados.extend(resultado.usados)
    if args.arquivar:
        destino = arquivar(resultado.usados, entrada, dia)
        print(f"Arquivos usados movidos para: {destino}")
    return 0


if __name__ == "__main__":
    try:
        sys.exit(main())
    except PermissionError as erro:
        print(f"\nERRO: não consegui gravar {erro.filename}.\n"
              "Se esse arquivo estiver aberto no Excel, feche-o e rode de novo.")
        sys.exit(1)
    except (FileNotFoundError, ValueError) as erro:
        print(f"\nERRO: {erro}")
        sys.exit(1)
